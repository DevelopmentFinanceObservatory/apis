"""Exports: a frame prepared once as a file a user keeps, behind a link that only signs.

`prepare_export` turns one exact frame version, a dataset's (`spec_`) or a derived frame's
(`ops_`), into a CSV, xlsx or Parquet file, alone or in a bundle zip, and records it under
`export_id`, a hash of the request alone. It resolves the version the request names and no
other. A request naming a dataset id alone (`BareExportFrame`) is first resolved to the version
the analysis runner reads for a Spec root of that id (`resolve_export_frame`), the newest in its
default enrichment, as the Presentation frames door reads a bare id. The request then runs, and
is keyed and answered, as one naming that pair. A dataset version's file is the one its file
selector serves (`doors.files`): the file the version lists under that rendering, else one
rendered from its stored frame. A derived frame's file is the one its file route renders
(`derived.render_derived_file`). The file is then copied to the pinned tier through
`retention_admission.copy_to_pinned`, which spends the caller's `pin` allowance and the day's
pin budgets as a pin does. A file already pinned spends nothing. The record keeps the pinned
copy's location and the facts the prepare answered.

A request whose id holds a record answers that record, rendering and spending nothing. Two
concurrent prepares of one request compute one id, and, since every render is deterministic
(bundle members carry fixed timestamps), one file, which the content-addressed store holds once:
the second copy finds the pinned bytes stored and gives its units back, and the second record
write returns the first's row.

The export link (`sign_export`) reads the record, refuses a version revoked since the prepare,
and signs a URL to the pinned copy. It never renders. The pinned tier never expires its objects,
so a link keeps working until its version is taken down.

The record keys the version by the id it is recorded under, the served id for a dataset the
request named by a posted id (`doors.snapshot.read_under_served_id`), as a pin's reference does.
The facts name the frame as the request did.

Refusals:

- A CSV or xlsx export of more than `EXPORT_MAX_ROWS` rows refuses `over_ceiling`
  (`ExportOverCeiling`). Parquet has no row ceiling.
- A derived frame stored without a recipe names no sources, and refuses
  `provenance_unavailable`.
- A source on the licence deny-list, `licence_restrictions.json` beside this module, keyed by
  source ref, refuses `licence_not_redistributable` at prepare and at the link. Every source is
  checked against the list as read when the process started, and the link checks the sources
  the record names, so a source added to the list withdraws its exports' links.
"""

from __future__ import annotations

import csv
import functools
import hashlib
import io
import json
import math
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any, Literal, cast

import openpyxl
import pyarrow.parquet as pq
import rfc8785
from pydantic import BaseModel, ConfigDict, Field

from artifact_store import ArtifactNotFoundError, ArtifactSigningError, Tier
from dataset_store.snapshot import SnapshotRevoked, read_servable_version
from datasets_api.admission import bounded_materialise_slot
from datasets_api.analysis import AnalysisPending
from datasets_api.analysis_roots import newest_spec_version
from datasets_api.artifact_delivery import artifact_content_disposition
from datasets_api.codes_exports import (
    ExportNotFound,
    ExportOverCeiling,
    LicenceNotRedistributable,
    NewestVersionUnservable,
    ProvenanceUnavailable,
)
from datasets_api.codes_snapshot import ArtifactExpired, SnapshotSigningFailed
from datasets_api.config import DOWNLOAD_SIGNED_URL_TTL
from datasets_api.current import DefinitionDeps
from datasets_api.deps import Deps
from datasets_api.derived import recipe_citation
from datasets_api.doors.derived import derived_file
from datasets_api.doors.files import (
    artifact_file_name,
    listed_file_location,
    read_version,
    render_file_location,
)
from datasets_api.doors.publish import PUBLISH_MAX_ROWS
from datasets_api.export_store import ExportRecord, ExportStore
from datasets_api.manifest import stored_manifest_facts
from datasets_api.manifest_models import CitationModel
from datasets_api.mint_models import MintOptions
from datasets_api.retention_admission import copy_to_pinned
from datasets_api.timings import phase
from dfo_platform.caller import Caller
from dfo_platform.codes import bound_registry_endpoint, registered_status
from dfo_platform.problem import Refusal
from dfo_platform.trace import get_trace_id
from materialise.formats import DATA_SHEET_TITLE

EXPORT_MAX_ROWS = PUBLISH_MAX_ROWS
"""The most rows a CSV or xlsx export holds: the frame ceiling a publish admits, 250,000."""

EXPORT_LINK_PREFIX = "/api/datasets/exports"
"""The path of the export link route, which the site's proxy serves on its own origin."""

_LICENCE_RESTRICTIONS = Path(__file__).with_name("licence_restrictions.json")

type ExportFormat = Literal["csv", "xlsx", "parquet"]


class ExportFrame(BaseModel):
    """One exact frame version: a dataset's `spec_` id or a derived frame's `ops_` id, and its
    content version."""

    model_config = ConfigDict(extra="forbid")
    spec_id: str = Field(pattern=r"^(spec|ops)_[^/]+$")
    content_version: str = Field(min_length=1)


class BareExportFrame(BaseModel):
    """A dataset named by its id alone, which stands for its newest version
    (`resolve_export_frame`)."""

    model_config = ConfigDict(extra="forbid")
    spec_id: str = Field(pattern=r"^spec_[0-9a-f]{32}$")


class ExportRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    frame: ExportFrame
    format: ExportFormat = Field(default="csv", description="the file format")
    bundle: bool = Field(
        default=False,
        description="true for a zip bundle of the file beside its README, datapackage.json and "
        "CITATION.cff",
    )


class ExportRequestBody(BaseModel):
    """`POST /api/datasets/exports`'s body: an `ExportRequest` whose frame may name a dataset id
    alone."""

    model_config = ConfigDict(extra="forbid")
    frame: ExportFrame | BareExportFrame = Field(
        description="the frame version to export, or `{spec_id}` alone for a dataset's newest "
        "version, which the answer's `frame` names"
    )
    format: ExportFormat = Field(default="csv", description="the file format")
    bundle: bool = Field(
        default=False,
        description="true for a zip bundle of the file beside its README, datapackage.json and "
        "CITATION.cff",
    )

    def resolved(self, frame: ExportFrame) -> ExportRequest:
        """The request of this body's format and bundle for `frame`."""
        return ExportRequest(frame=frame, format=self.format, bundle=self.bundle)


class ExportLicence(BaseModel):
    model_config = ConfigDict(extra="forbid")
    source: str = Field(description="the source's ref, the key of the licence deny-list")
    label: str
    licence: str | None = Field(description="the licence's label; null where none is declared")
    url: str | None = Field(description="the licence's terms page; null where none is declared")


class ExportResponse(BaseModel):
    """The facts of a prepared export. A repeat prepare answers the same facts."""

    model_config = ConfigDict(extra="forbid")
    export_id: str = Field(pattern=r"^exp_[0-9a-f]{32}$")
    url_path: str = Field(
        description="the export link's path, relative to the public site's origin. It needs no "
        "key and does not expire: each request answers 302 to a URL signed for five minutes"
    )
    file_name: str
    media_type: str
    format: ExportFormat
    bundle: bool
    bytes: int
    rows: int = Field(description="rows in the file's data table, header excluded")
    columns: int = Field(description="columns in the file's data table")
    sha256: str = Field(description="the hex SHA-256 of the file's bytes")
    frame: ExportFrame
    citation: str
    licences: list[ExportLicence] = Field(description="each source of the frame once")
    pinned: bool = Field(description="true: the file is held in the pinned tier, kept for good")


def export_id(request: ExportRequest) -> str:
    """`exp_` and the first 32 hex digits of the SHA-256 of the RFC 8785 form of the request's
    `{spec_id, content_version, format, bundle}`. A bare request is keyed by the pair it
    resolved to, so it shares its export with a request naming that pair."""
    preimage = {
        "spec_id": request.frame.spec_id,
        "content_version": request.frame.content_version,
        "format": request.format,
        "bundle": request.bundle,
    }
    return "exp_" + hashlib.sha256(rfc8785.dumps(preimage)).hexdigest()[:32]


@functools.cache
def restricted_sources() -> Mapping[str, str]:
    """The licence deny-list: each source ref whose licence does not allow redistribution, with
    the reason."""
    return cast(Mapping[str, str], json.loads(_LICENCE_RESTRICTIONS.read_text(encoding="utf-8")))


def check_licences(refs: Iterable[str]) -> None:
    """Refuse `LicenceNotRedistributable` when the deny-list names any of `refs`."""
    restricted = restricted_sources()
    listed = [(ref, restricted[ref]) for ref in dict.fromkeys(refs) if ref in restricted]
    if listed:
        raise LicenceNotRedistributable(sources=listed)


def _licences(citation: CitationModel) -> list[ExportLicence]:
    return [
        ExportLicence(
            source=source.ref, label=source.label, licence=source.license, url=source.license_url
        )
        for source in citation.sources_by_ref().values()
    ]


def file_shape(data: bytes, *, format: ExportFormat, bundle: bool) -> tuple[int, int]:
    """The rows, header excluded, and columns of the data table of the file `data`: a bundle's
    first data resource, and an xlsx workbook's `data` sheet."""
    if bundle:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            package = json.loads(archive.read("datapackage.json"))
            data = archive.read(package["resources"][0]["path"])
    if format == "parquet":
        metadata = pq.read_metadata(io.BytesIO(data))
        return metadata.num_rows, metadata.num_columns
    if format == "csv":
        reader = csv.reader(io.StringIO(data.decode("utf-8"), newline=""))
        header = next(reader)
        return sum(1 for _ in reader), len(header)
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    try:
        # `materialise.formats.write_excel` writes each sheet's dimension, so the sheet's size
        # is read from it in place of scanning every row.
        sheet = workbook[DATA_SHEET_TITLE]
        return cast(int, sheet.max_row) - 1, cast(int, sheet.max_column)
    finally:
        workbook.close()


def _check_rows(request: ExportRequest, rows: int) -> None:
    if request.format != "parquet" and rows > EXPORT_MAX_ROWS:
        raise ExportOverCeiling(format=request.format, count=rows, cap=EXPORT_MAX_ROWS)


def _pinned_copy(
    *,
    digest: str,
    tier: str,
    bucket: str,
    object_name: str,
    frame: ExportFrame,
    caller: Caller | None,
    current: Deps,
) -> tuple[str, str]:
    """The bucket and object name of the pinned copy of `digest`, stored at `tier`, copying it
    from the working tier (`copy_to_pinned`) unless it is pinned already. Working bytes gone
    before the copy refuse `ArtifactExpired`."""
    if tier == Tier.pinned:
        return bucket, object_name
    try:
        with phase("copy"):
            copied = copy_to_pinned(
                current.artifact_store,
                digest,
                allowances=current.allowances,
                caller=caller,
                now=current.now(),
            )
    except ArtifactNotFoundError as exc:
        raise ArtifactExpired(id=frame.spec_id, content_version=frame.content_version) from exc
    return copied.bucket, copied.object_name


def _prepare_dataset(
    request: ExportRequest,
    *,
    export: str,
    caller: Caller | None,
    current: Deps,
    deps: DefinitionDeps,
) -> ExportRecord:
    frame = request.frame
    stored, subject_ids = read_version(
        id=frame.spec_id, content_version=frame.content_version, deps=deps
    )
    facts = stored_manifest_facts(stored.version)
    citation = CitationModel.model_validate(facts.citation)
    licences = _licences(citation)
    check_licences(licence.source for licence in licences)
    _check_rows(request, cast(int, facts.coverage["rows"]))

    requested = MintOptions(format=request.format, bundle=request.bundle)
    location = listed_file_location(
        stored,
        id=frame.spec_id,
        requested=requested,
        enrichment_digest=deps.code_index.enrichment_digest,
        deps=deps,
    )
    if location is None:
        stored, location = render_file_location(
            stored, id=frame.spec_id, subject_ids=subject_ids, requested=requested, deps=deps
        )
    bucket, object_name = _pinned_copy(
        digest=location.artifact_digest,
        tier=location.tier,
        bucket=location.bucket,
        object_name=location.object_name,
        frame=frame,
        caller=caller,
        current=current,
    )
    with bounded_materialise_slot(
        current.materialise_semaphore,
        wait_seconds=current.settings.materialise_wait_seconds,
        deadline=deps.deadline,
    ):
        data = current.artifact_store.get(
            artifact_digest=location.artifact_digest, tier=Tier.pinned
        )
        rows, columns = file_shape(data, format=request.format, bundle=request.bundle)
    return _record(
        request,
        export=export,
        subject_id=stored.version.subject_id,
        digest=location.artifact_digest,
        bucket=bucket,
        object_name=object_name,
        file_name=artifact_file_name(location, stored=stored),
        media_type=location.media_type,
        size_bytes=location.size_bytes,
        rows=rows,
        columns=columns,
        citation=citation.text,
        licences=licences,
    )


def _prepare_derived(
    request: ExportRequest, *, export: str, caller: Caller | None, current: Deps
) -> ExportRecord:
    frame = request.frame
    read_servable_version(
        current.snapshot_reader,
        subject_id=frame.spec_id,
        content_version=frame.content_version,
        now=current.now(),
    )
    recipe = current.recipe_store.read_recipe(
        spec_id=frame.spec_id, content_version=frame.content_version
    )
    if recipe is None:
        raise ProvenanceUnavailable(spec_id=frame.spec_id, content_version=frame.content_version)
    citation = recipe_citation(current.snapshot_reader, recipe)
    licences = _licences(citation)
    check_licences(licence.source for licence in licences)

    rendered = derived_file(
        id=frame.spec_id,
        content_version=frame.content_version,
        options=MintOptions(format=request.format, bundle=request.bundle),
        current=current,
    )
    rows, columns = file_shape(rendered.data, format=request.format, bundle=request.bundle)
    _check_rows(request, rows)
    with phase("store"):
        stored = current.artifact_store.put(rendered.data, tier=Tier.working)
    bucket, object_name = _pinned_copy(
        digest=stored.digest,
        tier=stored.tier,
        bucket=stored.bucket,
        object_name=stored.object_name,
        frame=frame,
        caller=caller,
        current=current,
    )
    return _record(
        request,
        export=export,
        subject_id=frame.spec_id,
        digest=stored.digest,
        bucket=bucket,
        object_name=object_name,
        file_name=rendered.file_name,
        media_type=rendered.media_type,
        size_bytes=stored.size_bytes,
        rows=rows,
        columns=columns,
        citation=citation.text,
        licences=licences,
    )


def _record(
    request: ExportRequest,
    *,
    export: str,
    subject_id: str,
    digest: str,
    bucket: str,
    object_name: str,
    file_name: str,
    media_type: str,
    size_bytes: int,
    rows: int,
    columns: int,
    citation: str,
    licences: Sequence[ExportLicence],
) -> ExportRecord:
    body = ExportResponse(
        export_id=export,
        url_path=f"{EXPORT_LINK_PREFIX}/{export}",
        file_name=file_name,
        media_type=media_type,
        format=request.format,
        bundle=request.bundle,
        bytes=size_bytes,
        rows=rows,
        columns=columns,
        sha256=digest.removeprefix("sha256:"),
        frame=request.frame,
        citation=citation,
        licences=list(licences),
        pinned=True,
    )
    return ExportRecord(
        export_id=export,
        subject_id=subject_id,
        content_version=request.frame.content_version,
        artifact_digest=digest,
        bucket=bucket,
        object_name=object_name,
        body=body.model_dump(mode="json"),
    )


def resolve_export_frame(
    frame: ExportFrame | BareExportFrame, *, current: Deps, deps: DefinitionDeps
) -> ExportFrame:
    """`frame` itself when it names its version, else the dataset version its bare `spec_id`
    reads, as the analysis runner reads a Spec root of it (`analysis_roots.newest_spec_version`)
    and the Presentation frames door reads a bare id.

    Raises:
        NewestVersionUnservable: The version is building, or the serve refused with a code
            `datasets.exports` does not register.
        Exception: A refusal of the serve whose code `datasets.exports` registers, unchanged.
    """
    if isinstance(frame, ExportFrame):
        return frame
    try:
        dataset = newest_spec_version(
            frame.spec_id, deps=deps, recipes=current.recipe_store, trace_id=get_trace_id()
        )
    except AnalysisPending as pending:
        raise NewestVersionUnservable(
            code="service_unavailable",
            retryable=True,
            detail=f"the newest version of {frame.spec_id} is building; retry",
            headers={"Retry-After": str(math.ceil(pending.retry_after_s))},
        ) from None
    except Exception as exc:
        raise _unservable(frame.spec_id, exc) from exc
    return ExportFrame(spec_id=dataset.spec_id, content_version=dataset.content_version)


def passes_through(exc: BaseException) -> bool:
    """Whether the exports door answers `exc`, a failure resolving a bare `spec_id`, unchanged:
    it is no refusal, so the surface classifies it, or a refusal whose code `datasets.exports`
    registers."""
    return not isinstance(exc, Refusal) or _status(exc.code, "datasets.exports") is not None


def _status(code: str, endpoint: str) -> int | None:
    """The status the code registry gives `code` on `endpoint`, `None` where it gives none."""
    try:
        with bound_registry_endpoint(endpoint):
            return registered_status(code)
    except LookupError:
        return None


def _unservable(spec_id: str, exc: Exception) -> Exception:
    """The refusal the exports door answers for `exc`, a failure resolving `spec_id`: `exc`
    itself when it `passes_through`, else a `NewestVersionUnservable` of the registered code
    that stands for it, read from the status the analysis door, whose resolve this is, gives
    its code. A code neither endpoint registers is the server's own rendering failing, an
    `internal_error`. Nothing is classified here, so the door classifies the failure once."""
    if passes_through(exc):
        return exc
    assert isinstance(exc, Refusal)
    detail = f"the newest version of {spec_id} cannot be served: {exc.code}: {exc.detail or exc}"
    status = _status(exc.code, "datasets.analysis")
    if status is None:
        return NewestVersionUnservable(
            code="internal_error",
            detail=f"the newest version of {spec_id} cannot be served: {exc}",
        )
    if exc.retryable:
        return NewestVersionUnservable(
            code="service_unavailable", retryable=True, detail=detail, headers=exc.headers
        )
    if status >= 500:
        return NewestVersionUnservable(code="internal_error", detail=detail)
    return NewestVersionUnservable(code="snapshot_not_found", detail=detail)


def recorded_export(*, request: ExportRequest, store: ExportStore) -> dict[str, Any] | None:
    """The facts of the export recorded under the request's id, or `None` when none is.

    Raises:
        LicenceNotRedistributable: The deny-list names a source of the export.
    """
    record = store.read_export(export_id(request))
    if record is None:
        return None
    return _servable_body(record)


def _servable_body(record: ExportRecord) -> dict[str, Any]:
    body = ExportResponse.model_validate(record.body)
    check_licences(licence.source for licence in body.licences)
    return body.model_dump(mode="json")


def prepare_export(
    *, request: ExportRequest, caller: Caller | None, current: Deps, deps: DefinitionDeps
) -> dict[str, Any]:
    """Prepare the file `request` names, pin it, record it, and return its facts. See the
    module docstring. `caller` names an open caller, whose `pin` allowance a new pinned copy
    spends, and `None` a caller holding the pin grant. `deps` is the request's retrieval view of
    `current`, under its deadline.

    Raises:
        SpecNotFound, IdentityQuarantined, SnapshotNotFound, SnapshotRevoked, ArtifactExpired:
            The version cannot be read, or its file can no longer be rendered.
        ProvenanceUnavailable: A derived frame has no stored recipe.
        LicenceNotRedistributable: The deny-list names a source of the frame.
        ExportOverCeiling: A CSV or xlsx export of more than `EXPORT_MAX_ROWS` rows.
        PinBudgetExhausted, OpenPinBudgetExhausted, AllowanceExhausted: The copy's units are
            spent.
    """
    export = export_id(request)
    if request.frame.spec_id.startswith("ops_"):
        record = _prepare_derived(request, export=export, caller=caller, current=current)
    else:
        record = _prepare_dataset(request, export=export, caller=caller, current=current, deps=deps)
    with phase("record"):
        stored = current.export_store.write_export(record)
    return _servable_body(stored)


def sign_export(*, export_id: str, current: Deps) -> str:
    """A URL to the export's pinned file, signed for `DOWNLOAD_SIGNED_URL_TTL` and naming the
    download by its file name.

    The version is read again first, so a takedown reaches the link as it reaches every other
    file link: a revocation of the version or its base version under the id it is recorded
    under, or under the id the request named, refuses.

    Raises:
        ExportNotFound: No export is recorded under `export_id`.
        SnapshotRevoked: The version, or its base version, is revoked.
        LicenceNotRedistributable: The deny-list names a source of the export.
        SnapshotSigningFailed: Signing failed.
    """
    record = current.export_store.read_export(export_id)
    if record is None:
        raise ExportNotFound(export_id=export_id)
    named = cast(Mapping[str, str], record.body["frame"])
    with phase("revocation"):
        state = current.snapshot_reader.read_version_state(
            record.subject_id, record.content_version, revoked_under=(named["spec_id"],)
        )
    # Version records are never deleted, and the export row's foreign key names one.
    assert state is not None
    if state.revocation is not None:
        raise SnapshotRevoked(
            id=named["spec_id"],
            content_version=record.content_version,
            reason=state.revocation.reason,
        )
    body = _servable_body(record)
    try:
        with phase("sign"):
            return current.artifact_store.signed_location_url(
                record.artifact_digest,
                tier=Tier.pinned,
                bucket=record.bucket,
                object_name=record.object_name,
                ttl=DOWNLOAD_SIGNED_URL_TTL,
                response_type=body["media_type"],
                response_disposition=artifact_content_disposition(body["file_name"]),
            )
    except ArtifactSigningError as exc:
        # Copy safe fields here, then raise outside the handler so Python leaves `__context__`
        # unset, as the files door does.
        failure = (exc.error_class_name, exc.status_code)
    raise SnapshotSigningFailed(
        bucket=record.bucket,
        object_name=record.object_name,
        digest=record.artifact_digest,
        error_class_name=failure[0],
        status_code=failure[1],
    )
