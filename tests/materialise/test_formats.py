"""Writer, schema, cap, and streaming tests for `materialise.formats`.
Excel and Parquet checks compare content and repeat-write digests because
library metadata affects bytes (`openpyxl==3.1.5`, `pyarrow==24.0.0`). CSV
has a byte-exact golden fixture.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import time
from dataclasses import dataclass
from datetime import datetime

import openpyxl
import pyarrow.parquet as pq
import pytest
from dfo_frame_grammar import Frame, canonical_frame_bytes
from dfo_frame_grammar.grammar import ComposedUnit, StatVarPart, UnitRequest

import materialise.formats as materialise_formats
from materialise.errors import MaterialisationError, MaterialisationUsageError
from materialise.formats import (
    EXCEL_MAX_CELL_CHARS,
    EXCEL_MAX_COLUMNS,
    FRAME1_MEDIA_TYPE,
    ExcelIllegalCharacterError,
    ExcelLimitError,
    ParquetSizeBoundError,
    Rendering,
    _sanitize_csv_cell,
    frame1_entity_columns,
    read_frame1_table,
    table_schema_for,
)
from materialise.formats import write_csv as _real_write_csv
from materialise.formats import write_excel as _real_write_excel
from materialise.formats import write_frame1 as _real_write_frame1
from materialise.formats import write_parquet as _real_write_parquet
from materialise.frame_build import build_frame, unit_cell_for_rebuild

from .conftest import read_fixture_text


class _Sink:
    """Minimal in-memory `materialise.formats.Sink`: buffers every byte
    written, so a test can inspect what actually reached it after the
    writer call returns.
    """

    def __init__(self) -> None:
        self._buffer = bytearray()

    def write(self, data: bytes, /) -> int:
        self._buffer += data
        return len(data)

    def getvalue(self) -> bytes:
        return bytes(self._buffer)


@dataclass(frozen=True)
class _Written:
    """Artifact metadata and bytes captured by its sink."""

    data: bytes
    media_type: str
    digest: str


def write_csv(chunks, *, rendering: Rendering) -> _Written:
    sink = _Sink()
    artifact = _real_write_csv(chunks, rendering=rendering, sink=sink)
    return _Written(data=sink.getvalue(), media_type=artifact.media_type, digest=artifact.digest)


def write_excel(chunks, *, rendering: Rendering) -> _Written:
    sink = _Sink()
    artifact = _real_write_excel(chunks, rendering=rendering, sink=sink)
    return _Written(data=sink.getvalue(), media_type=artifact.media_type, digest=artifact.digest)


def write_parquet(chunks, *, rendering: Rendering, size_bound: int | None = None) -> _Written:
    sink = _Sink()
    kwargs = {} if size_bound is None else {"size_bound": size_bound}
    artifact = _real_write_parquet(chunks, rendering=rendering, sink=sink, **kwargs)
    return _Written(data=sink.getvalue(), media_type=artifact.media_type, digest=artifact.digest)


def write_frame1(chunks) -> _Written:
    sink = _Sink()
    artifact = _real_write_frame1(chunks, sink=sink)
    assert artifact.size_bytes == len(sink.getvalue())
    return _Written(data=sink.getvalue(), media_type=artifact.media_type, digest=artifact.digest)


@dataclass(frozen=True)
class _Table:
    """A plain literal table satisfying `formats.Table`'s structural
    contract.
    """

    columns: tuple[str, ...]
    rows: tuple[tuple[object, ...], ...]


def _frame1_source_frame():
    """Build a real frame for the `write_frame1` tests."""
    return build_frame(
        role_keys=("recipient",),
        rows=[
            {
                "stat_var": "dcid:MalariaGrantDisbursements",
                "recipient": "country/TGO",
                "time_period": "2015",
                "value": 1000.0,
                "unit": "USD",
                "measurementMethod": "GrantEquivalent",
                "provenance": "OECD_DAC2A",
            },
            {
                "stat_var": "dcid:MalariaGrantDisbursements",
                "recipient": "country/KEN",
                "time_period": "2016",
                "value": 2500.5,
                "unit": None,
                "measurementMethod": None,
                "provenance": None,
            },
        ],
    )


_GOLDEN_TABLE = _Table(
    columns=("stat_var", "recipient", "time_period", "value", "unit"),
    rows=(
        ("dcid:MalariaGrantDisbursements", "country/TGO", "2015", 1000.0, "USD"),
        ("dcid:MalariaGrantDisbursements", "country/KEN", "2016", 2500.5, None),
    ),
)

_FRIENDLY_RENDERING = Rendering(
    friendly_headers={
        "stat_var": "Statistical Variable",
        "recipient": "Recipient",
        "time_period": "Period",
        "value": "Value",
        "unit": "Unit",
    }
)

_PLAIN_RENDERING = Rendering()

_INJECTION_TABLE = _Table(
    columns=("label",),
    rows=(
        ("=cmd|' /C calc'!A1",),
        ("+1+1",),
        ("-2+3",),
        ("@SUM(A1:A9)",),
        ("\ttab-lead",),
        ("\rcr-lead",),
        ("plain text",),
        ("'already-quoted",),
        (None,),
        (42,),
    ),
)


_TRIGGER_HEADER_TABLE = _Table(columns=("col",), rows=(("value",),))
_TRIGGER_HEADER_RENDERING = Rendering(friendly_headers={"col": "=cmd"})

_DUPLICATE_HEADER_TABLE = _Table(columns=("a", "b"), rows=(("1", "2"),))
_DUPLICATE_HEADER_RENDERING = Rendering(friendly_headers={"a": "same", "b": "same"})

_RAW_DUPLICATE_COLUMNS_TABLE = _Table(columns=("dup", "dup"), rows=(("1", "2"),))


class _Stringable:
    """A tiny object with a deterministic `__str__`, standing in for an
    arbitrary non-scalar cell (e.g. a domain object CSV/Parquet already
    stringify)."""

    def __init__(self, text: str) -> None:
        self._text = text

    def __str__(self) -> str:
        return self._text


def _recompute_digest(data: bytes) -> str:
    return "sha256:" + hashlib.sha256(data).hexdigest()


# --- write_csv -------------------------------------------------------------


def test_write_csv_matches_checked_in_golden_bytes():
    artifact = write_csv((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    expected = read_fixture_text("formats-golden.csv").encode("utf-8")
    assert artifact.data == expected
    assert artifact.media_type == "text/csv"


def test_write_csv_digest_matches_recomputed_sha256():
    artifact = write_csv((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    assert artifact.digest == _recompute_digest(artifact.data)


def test_write_csv_without_friendly_headers_uses_raw_column_names():
    artifact = write_csv((_GOLDEN_TABLE,), rendering=_PLAIN_RENDERING)
    header_line = artifact.data.split(b"\r\n", 1)[0]
    assert header_line == b"stat_var,recipient,time_period,value,unit"


def test_write_csv_neutralises_formula_trigger_cells():
    artifact = write_csv((_INJECTION_TABLE,), rendering=_PLAIN_RENDERING)
    text = artifact.data.decode("utf-8")
    # CSV quoting of embedded CR and tab is separate from formula neutralisation.
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0] == ["label"]
    assert rows[1] == ["'=cmd|' /C calc'!A1"]
    assert rows[2] == ["'+1+1"]
    assert rows[3] == ["'-2+3"]
    assert rows[4] == ["'@SUM(A1:A9)"]
    assert rows[5] == ["'\ttab-lead"]
    assert rows[6] == ["'\rcr-lead"]
    assert rows[7] == ["plain text"]
    assert rows[8] == ["'already-quoted"]
    assert rows[9] == [""]  # None renders as an empty field
    assert rows[10] == ["42"]


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("=cmd", "'=cmd"),
        ("+1", "'+1"),
        ("-1", "'-1"),
        ("@SUM(A1)", "'@SUM(A1)"),
        ("\ttab", "'\ttab"),
        ("\rcr", "'\rcr"),
        ("plain", "plain"),
        ("'already", "'already"),
        # Trigger cells with leading apostrophes gain one more apostrophe,
        # preserving distinct raw values.
        ("'=foo", "''=foo"),
        ("''=foo", "'''=foo"),
        (None, None),
        (42, 42),
        (3.5, 3.5),
    ],
)
def test_sanitize_csv_cell(value, expected):
    assert _sanitize_csv_cell(value) == expected


def test_sanitize_csv_cell_collision_pair_round_trips_to_distinct_originals():
    # Counting leading apostrophes keeps sanitised `=foo` and `'=foo`
    # distinct and reversible.
    bare_written = _sanitize_csv_cell("=foo")
    quoted_written = _sanitize_csv_cell("'=foo")
    assert bare_written == "'=foo"
    assert quoted_written == "''=foo"
    assert bare_written != quoted_written


def test_sanitize_csv_cell_is_reversible_for_neutralised_cells():
    # Removing one added apostrophe must recover each original trigger cell.
    for original in (
        "=cmd",
        "+1+1",
        "-2+3",
        "@SUM(A1)",
        "\ttab",
        "\rcr",
        "'=cmd",
        "''=cmd",
    ):
        neutralised = _sanitize_csv_cell(original)
        assert neutralised.startswith("'")
        stripped = neutralised[1:]
        assert stripped == original


# --- table_schema_for --------------------------------------------------------


def test_table_schema_for_golden_table():
    schema = table_schema_for(_GOLDEN_TABLE, rendering=_FRIENDLY_RENDERING)
    assert schema == {
        "fields": [
            {"name": "Statistical Variable", "type": "string"},
            {"name": "Recipient", "type": "string"},
            {"name": "Period", "type": "string"},
            {"name": "Value", "type": "number"},
            {"name": "Unit", "type": "string"},
        ]
    }


def test_table_schema_for_infers_types_from_cells_not_headers():
    table = _Table(
        columns=("ints", "bools", "mixed_numeric", "all_null", "mixed_incompatible"),
        rows=(
            (1, True, 1, None, "x"),
            (2, False, 2.5, None, 1),
        ),
    )
    schema = table_schema_for(table, rendering=_PLAIN_RENDERING)
    types = {field["name"]: field["type"] for field in schema["fields"]}
    assert types == {
        "ints": "integer",
        "bools": "boolean",
        "mixed_numeric": "number",
        "all_null": "any",
        "mixed_incompatible": "any",
    }


# All four writers use sanitised, unique headers.


def test_a_trigger_leading_friendly_header_is_sanitized_identically_everywhere():
    csv_artifact = write_csv((_TRIGGER_HEADER_TABLE,), rendering=_TRIGGER_HEADER_RENDERING)
    csv_header_line = csv_artifact.data.split(b"\r\n", 1)[0].decode("utf-8")
    assert csv_header_line == "'=cmd"

    excel_artifact = write_excel((_TRIGGER_HEADER_TABLE,), rendering=_TRIGGER_HEADER_RENDERING)
    workbook = openpyxl.load_workbook(io.BytesIO(excel_artifact.data))
    excel_header = next(workbook.active.iter_rows(max_row=1))[0].value
    assert excel_header == "'=cmd"

    parquet_artifact = write_parquet((_TRIGGER_HEADER_TABLE,), rendering=_TRIGGER_HEADER_RENDERING)
    parquet_table = pq.read_table(io.BytesIO(parquet_artifact.data))
    assert parquet_table.schema.names == ["'=cmd"]

    schema = table_schema_for(_TRIGGER_HEADER_TABLE, rendering=_TRIGGER_HEADER_RENDERING)
    assert schema["fields"][0]["name"] == "'=cmd"


@pytest.mark.parametrize(
    "call",
    [
        lambda: write_csv((_DUPLICATE_HEADER_TABLE,), rendering=_DUPLICATE_HEADER_RENDERING),
        lambda: write_excel((_DUPLICATE_HEADER_TABLE,), rendering=_DUPLICATE_HEADER_RENDERING),
        lambda: write_parquet((_DUPLICATE_HEADER_TABLE,), rendering=_DUPLICATE_HEADER_RENDERING),
        lambda: table_schema_for(_DUPLICATE_HEADER_TABLE, rendering=_DUPLICATE_HEADER_RENDERING),
    ],
    ids=["write_csv", "write_excel", "write_parquet", "table_schema_for"],
)
def test_friendly_headers_mapping_two_columns_to_one_label_refuses_everywhere(call):
    with pytest.raises(MaterialisationUsageError) as exc_info:
        call()
    assert "same" in str(exc_info.value)


def test_a_tables_own_duplicate_column_names_refuse_even_without_friendly_headers():
    with pytest.raises(MaterialisationUsageError) as exc_info:
        table_schema_for(_RAW_DUPLICATE_COLUMNS_TABLE, rendering=_PLAIN_RENDERING)
    assert "dup" in str(exc_info.value)


# Excel content, formula safety, and byte stability.


@pytest.mark.parametrize("read_only", [False, True], ids=["loaded", "streamed"])
def test_write_excel_content_round_trips(read_only):
    """A streaming read, as frictionless makes, sizes rows by the sheet's dimension, so a row
    ending in an empty cell reads at full width only when the writer records that range."""
    artifact = write_excel((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    assert (
        artifact.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = openpyxl.load_workbook(io.BytesIO(artifact.data), read_only=read_only)
    sheet = workbook.active
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert rows[0] == ["Statistical Variable", "Recipient", "Period", "Value", "Unit"]
    assert rows[1] == ["dcid:MalariaGrantDisbursements", "country/TGO", "2015", 1000, "USD"]
    assert rows[2] == ["dcid:MalariaGrantDisbursements", "country/KEN", "2016", 2500.5, None]


def test_write_excel_neutralises_formula_trigger_cells_and_they_are_inert():
    artifact = write_excel((_INJECTION_TABLE,), rendering=_PLAIN_RENDERING)
    workbook = openpyxl.load_workbook(io.BytesIO(artifact.data))
    sheet = workbook.active
    cells = [row[0] for row in sheet.iter_rows(min_row=2)]
    # Formula-trigger cells load as strings.
    assert all(cell.data_type != "f" for cell in cells)
    values = [cell.value for cell in cells]
    assert values == [
        "'=cmd|' /C calc'!A1",
        "'+1+1",
        "'-2+3",
        "'@SUM(A1:A9)",
        "'\ttab-lead",
        # XML normalizes bare CR to LF when parsing text.
        "'\ncr-lead",
        "plain text",
        "'already-quoted",
        None,
        42,
    ]


def test_write_excel_digest_matches_recomputed_sha256():
    artifact = write_excel((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    assert artifact.digest == _recompute_digest(artifact.data)


def test_write_excel_normalizes_an_arbitrary_object_cell_to_its_str():
    table = _Table(columns=("only",), rows=((_Stringable("custom-object"),),))
    artifact = write_excel((table,), rendering=_PLAIN_RENDERING)
    workbook = openpyxl.load_workbook(io.BytesIO(artifact.data))
    cell = next(workbook.active.iter_rows(min_row=2))[0]
    assert cell.value == "custom-object"


def test_write_excel_normalizes_then_sanitizes_an_object_whose_str_starts_with_a_trigger():
    table = _Table(columns=("only",), rows=((_Stringable("=cmd"),),))
    artifact = write_excel((table,), rendering=_PLAIN_RENDERING)
    workbook = openpyxl.load_workbook(io.BytesIO(artifact.data))
    cell = next(workbook.active.iter_rows(min_row=2))[0]
    assert cell.value == "'=cmd"
    assert cell.data_type != "f"


def test_write_excel_normalizes_a_datetime_cell_the_same_way_csv_and_parquet_do():
    when = datetime(2024, 1, 1, 12, 30)
    table = _Table(columns=("only",), rows=((when,),))

    excel_artifact = write_excel((table,), rendering=_PLAIN_RENDERING)
    workbook = openpyxl.load_workbook(io.BytesIO(excel_artifact.data))
    excel_cell = next(workbook.active.iter_rows(min_row=2))[0].value

    csv_artifact = write_csv((table,), rendering=_PLAIN_RENDERING)
    csv_row = list(csv.reader(io.StringIO(csv_artifact.data.decode("utf-8"))))[1]

    assert excel_cell == str(when)
    assert csv_row[0] == str(when)


@pytest.mark.parametrize(
    ("unit", "expected"),
    [
        (
            UnitRequest(currency="EUR", prices="constant", base_year=2024),
            '{"base_year":2024,"currency":"EUR","prices":"constant"}',
        ),
        (
            ComposedUnit(
                numerator=UnitRequest(currency="EUR", prices="current"),
                denominator=StatVarPart(stat_var="dcid:Count_Person"),
            ),
            '{"denominator":{"stat_var":"dcid:Count_Person"},'
            '"numerator":{"currency":"EUR","prices":"current"}}',
        ),
    ],
    ids=["converted", "composed"],
)
def test_every_tabular_writer_renders_a_structured_unit_cell_as_canonical_json(unit, expected):
    table = _Table(columns=("unit",), rows=((unit,),))

    csv_artifact = write_csv((table,), rendering=_PLAIN_RENDERING)
    excel_artifact = write_excel((table,), rendering=_PLAIN_RENDERING)
    parquet_artifact = write_parquet((table,), rendering=_PLAIN_RENDERING)

    csv_row = list(csv.reader(io.StringIO(csv_artifact.data.decode("utf-8"))))[1]
    workbook = openpyxl.load_workbook(io.BytesIO(excel_artifact.data))
    excel_cell = next(workbook.active.iter_rows(min_row=2))[0].value
    parquet_cell = pq.read_table(io.BytesIO(parquet_artifact.data)).column(0).to_pylist()[0]
    assert csv_row == [expected]
    assert excel_cell == expected
    assert parquet_cell == expected


def test_write_excel_is_byte_stable_across_a_real_wall_clock_second_boundary():
    # openpyxl timestamps properties and ZIP members to the second. The
    # 1.1-second sleep crosses a timestamp boundary.
    first = write_excel((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    time.sleep(1.1)
    second = write_excel((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    assert first.digest == second.digest
    assert first.data == second.data


# Excel row, column, and cell-length limits.


def test_write_excel_refuses_over_column_cap():
    table = _Table(
        columns=tuple(f"c{i}" for i in range(EXCEL_MAX_COLUMNS + 1)),
        rows=(),
    )
    with pytest.raises(ExcelLimitError) as exc_info:
        write_excel((table,), rendering=_PLAIN_RENDERING)
    error = exc_info.value
    assert isinstance(error, MaterialisationError)
    assert error.code == "excel_limit"
    assert error.status == 422
    assert str(EXCEL_MAX_COLUMNS + 1) in str(error)
    assert "csv" in str(error).lower() or "parquet" in str(error).lower()


def test_write_excel_refuses_over_row_cap(monkeypatch):
    # Two data rows plus the header exceed a two-row cap by one.
    row_cap = 2
    monkeypatch.setattr(materialise_formats, "EXCEL_MAX_ROWS", row_cap)
    table = _Table(columns=("only",), rows=tuple((None,) for _ in range(row_cap)))
    with pytest.raises(ExcelLimitError) as exc_info:
        write_excel((table,), rendering=_PLAIN_RENDERING)
    error = exc_info.value
    assert error.code == "excel_limit"
    assert str(row_cap + 1) in str(error)


def test_write_excel_refuses_over_cell_character_cap():
    table = _Table(columns=("only",), rows=(("x" * (EXCEL_MAX_CELL_CHARS + 1),),))
    with pytest.raises(ExcelLimitError) as exc_info:
        write_excel((table,), rendering=_PLAIN_RENDERING)
    error = exc_info.value
    assert error.code == "excel_limit"
    assert str(EXCEL_MAX_CELL_CHARS + 1) in str(error)


def test_write_excel_within_all_three_caps_succeeds():
    table = _Table(columns=("only",), rows=(("fits",),))
    artifact = write_excel((table,), rendering=_PLAIN_RENDERING)
    assert artifact.data


def test_write_excel_refuses_when_a_trigger_leading_cell_at_the_raw_cap_would_overflow_on_write():
    # Sanitisation adds one quote to a trigger cell at the cap. Check the
    # written value to avoid openpyxl truncation.
    raw = "=" + "x" * (EXCEL_MAX_CELL_CHARS - 1)
    assert len(raw) == EXCEL_MAX_CELL_CHARS
    table = _Table(columns=("only",), rows=((raw,),))
    with pytest.raises(ExcelLimitError) as exc_info:
        write_excel((table,), rendering=_PLAIN_RENDERING)
    error = exc_info.value
    assert error.code == "excel_limit"
    assert str(EXCEL_MAX_CELL_CHARS + 1) in str(error)


def test_write_excel_non_trigger_twin_at_the_exact_cell_cap_still_succeeds():
    # A non-trigger cell at the cap needs no extra sanitisation character.
    raw = "x" * EXCEL_MAX_CELL_CHARS
    table = _Table(columns=("only",), rows=((raw,),))
    artifact = write_excel((table,), rendering=_PLAIN_RENDERING)
    assert artifact.data


# Excel illegal-character handling.


def test_write_excel_refuses_a_cell_with_an_xml_illegal_control_character():
    table = _Table(columns=("a", "b"), rows=(("fine", "bad\x01value"),))
    with pytest.raises(ExcelIllegalCharacterError) as exc_info:
        write_excel((table,), rendering=_PLAIN_RENDERING)
    error = exc_info.value
    assert isinstance(error, MaterialisationError)
    assert error.code == "excel_illegal_character"
    assert error.status == 422
    # The header is row 0, so the one data row is row 1. Column 1 is "b".
    assert error.row == 1
    assert error.column == 1


def test_write_excel_refuses_an_illegal_character_in_the_header_row():
    table = _Table(columns=("fine", "bad\x02header"), rows=(("x", "y"),))
    with pytest.raises(ExcelIllegalCharacterError) as exc_info:
        write_excel((table,), rendering=_PLAIN_RENDERING)
    error = exc_info.value
    assert error.row == 0
    assert error.column == 1


def test_write_excel_tolerates_tab_and_newline_which_stay_xml_legal():
    # `\t` and `\n` are XML-legal whitespace and pass through unchanged.
    table = _Table(columns=("only",), rows=(("line one\nline two\ttabbed",),))
    artifact = write_excel((table,), rendering=_PLAIN_RENDERING)
    assert artifact.data


def test_write_excel_a_clean_frame_with_no_illegal_characters_renders():
    table = _Table(columns=("a", "b"), rows=(("clean", "also clean"), ("more", "text")))
    artifact = write_excel((table,), rendering=_PLAIN_RENDERING)
    assert artifact.data


# --- write_parquet -----------------------------------------------------------


def test_write_parquet_content_round_trips():
    artifact = write_parquet((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    assert artifact.media_type == "application/vnd.apache.parquet"
    table = pq.read_table(io.BytesIO(artifact.data))
    assert table.to_pydict() == {
        "Statistical Variable": [
            "dcid:MalariaGrantDisbursements",
            "dcid:MalariaGrantDisbursements",
        ],
        "Recipient": ["country/TGO", "country/KEN"],
        "Period": ["2015", "2016"],
        "Value": [1000.0, 2500.5],
        "Unit": ["USD", None],
    }
    assert [f.name for f in table.schema] == [
        "Statistical Variable",
        "Recipient",
        "Period",
        "Value",
        "Unit",
    ]


def test_write_parquet_digest_matches_recomputed_sha256():
    artifact = write_parquet((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    assert artifact.digest == _recompute_digest(artifact.data)


def test_write_parquet_is_byte_stable_across_a_real_wall_clock_second_boundary():
    first = write_parquet((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    time.sleep(1.1)
    second = write_parquet((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    assert first.digest == second.digest
    assert first.data == second.data


def test_write_parquet_refuses_over_its_size_bound():
    # A small table with a tiny injected size_bound exercises the refusal
    # without needing to materialise a genuinely huge table.
    table = _Table(columns=("value",), rows=tuple((f"row-{i}",) for i in range(500)))
    with pytest.raises(ParquetSizeBoundError) as exc_info:
        write_parquet((table,), rendering=_PLAIN_RENDERING, size_bound=64)
    error = exc_info.value
    assert isinstance(error, MaterialisationError)
    assert error.code == "parquet_size_bound"
    assert error.status == 422
    assert "64" in str(error)


def test_write_parquet_within_its_size_bound_succeeds():
    table = _Table(columns=("value",), rows=(("fits",),))
    artifact = write_parquet((table,), rendering=_PLAIN_RENDERING, size_bound=1_000_000)
    assert artifact.data


def test_write_parquet_refuses_on_an_empty_table_whose_batch_loop_never_runs():
    # An empty table has no batches, so only the closing-time check sees bound 0.
    table = _Table(columns=("value",), rows=())
    with pytest.raises(ParquetSizeBoundError) as exc_info:
        write_parquet((table,), rendering=_PLAIN_RENDERING, size_bound=0)
    error = exc_info.value
    assert error.code == "parquet_size_bound"
    assert "0" in str(error)


def test_write_parquet_refuses_when_only_the_closing_footer_crosses_the_bound():
    # `writer.close()` adds a footer after the in-loop size check. The final
    # closed size must obey the bound.
    table = _Table(columns=("value",), rows=(("fits",),))
    baseline = write_parquet((table,), rendering=_PLAIN_RENDERING)
    bound = len(baseline.data) - 1
    with pytest.raises(ParquetSizeBoundError) as exc_info:
        write_parquet((table,), rendering=_PLAIN_RENDERING, size_bound=bound)
    error = exc_info.value
    assert error.code == "parquet_size_bound"
    assert str(bound) in str(error)


# --- write_frame1 -----------------------------------------------------


def test_write_frame1_bytes_are_exactly_canonical_frame_bytes():
    frame = _frame1_source_frame()
    artifact = write_frame1((frame,))
    assert artifact.data == canonical_frame_bytes(frame)


def test_write_frame1_media_type_is_the_published_constant():
    artifact = write_frame1((_frame1_source_frame(),))
    assert artifact.media_type == FRAME1_MEDIA_TYPE == "application/vnd.dfo.frame1+json"


def test_write_frame1_digest_matches_recomputed_sha256():
    artifact = write_frame1((_frame1_source_frame(),))
    assert artifact.digest == _recompute_digest(artifact.data)


def test_write_frame1_is_deterministic_across_two_calls_over_the_same_frame():
    frame = _frame1_source_frame()
    first = write_frame1((frame,))
    second = write_frame1((frame,))
    assert first.data == second.data
    assert first.digest == second.digest


def _many_row_frame(rows: int):
    return build_frame(
        role_keys=("provider", "recipient"),
        rows=[
            {
                "stat_var": "dcid:Amount_FinancialFlow_Gross_Disbursement",
                "provider": f"country/P{index % 7:02d}",
                "recipient": f"country/R{index:05d}",
                "time_period": str(2000 + index % 20),
                "value": float(index),
                "unit": "USD",
                "measurementMethod": None,
                "provenance": "OECD_CRS",
            }
            for index in range(rows)
        ],
    )


def _chunks_of(frame, size: int):
    return [
        Frame(columns=frame.columns, table=frame.table.slice(start, size))
        for start in range(0, frame.table.num_rows, size)
    ]


def test_write_frame1_over_several_chunks_gives_the_bytes_and_digest_of_one_whole_frame():
    frame = _many_row_frame(50)
    chunks = _chunks_of(frame, 7)
    assert len(chunks) == 8

    whole = canonical_frame_bytes(frame)
    artifact = write_frame1(chunks)

    assert artifact.data == whole
    assert artifact.digest == _recompute_digest(whole)


def test_write_frame1_streams_a_multi_chunk_frame_in_bounded_writes(
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setattr(materialise_formats, "_FRAME1_WRITE_BYTES", 512)
    frame = _many_row_frame(200)
    whole = canonical_frame_bytes(frame)

    class _RecordingSink:
        def __init__(self) -> None:
            self.writes: list[int] = []

        def write(self, data: bytes, /) -> int:
            self.writes.append(len(data))
            return len(data)

    sink = _RecordingSink()
    artifact = _real_write_frame1(_chunks_of(frame, 32), sink=sink)

    assert artifact.size_bytes == sum(sink.writes) == len(whole)
    assert len(sink.writes) > 10
    # Each write holds one block plus at most one row's bytes, far below the document size.
    assert max(sink.writes) < 1024 < len(whole)


# --- read_frame1_table / frame1_entity_columns -----------------------------


def test_read_frame1_table_round_trips_write_csv_bytes():
    # The frame contains `1000.0`; JSON parsing can turn it into an int.
    # CSV equality confirms `read_frame1_table` restores the float.
    frame = _frame1_source_frame()
    table = read_frame1_table(write_frame1((frame,)).data)
    assert (
        write_csv((frame,), rendering=_PLAIN_RENDERING).data
        == write_csv((table,), rendering=_PLAIN_RENDERING).data
    )


def test_read_frame1_table_round_trips_write_parquet_bytes():
    # Equal Parquet bytes confirm the read-back table preserves rows, order,
    # and the whole-number value cell.
    frame = _frame1_source_frame()
    table = read_frame1_table(write_frame1((frame,)).data)
    assert (
        write_parquet((frame,), rendering=_PLAIN_RENDERING).data
        == write_parquet((table,), rendering=_PLAIN_RENDERING).data
    )


@pytest.mark.parametrize(
    ("unit", "cell"),
    [
        (
            {"currency": "EUR", "prices": "constant", "base_year": 2024},
            UnitRequest(currency="EUR", prices="constant", base_year=2024),
        ),
        (
            {"numerator": "USDollar", "denominator": {"stat_var": "dcid:Count_Person"}},
            ComposedUnit(
                numerator="USDollar", denominator=StatVarPart(stat_var="dcid:Count_Person")
            ),
        ),
    ],
    ids=["converted", "composed"],
)
def test_read_frame1_table_types_a_structured_unit_and_its_rows_rebuild_the_frame(unit, cell):
    role_keys = ("recipient",)
    frame = build_frame(
        role_keys=role_keys,
        rows=[
            {
                "stat_var": "dcid:Exemplar",
                "recipient": "country/TGO",
                "time_period": "2015",
                "value": 1000.0,
                "unit": unit,
            }
        ],
    )
    table = read_frame1_table(write_frame1((frame,)).data)

    assert [row[table.columns.index("unit")] for row in table.rows] == [cell]
    rebuilt = build_frame(
        role_keys=role_keys,
        rows=[
            {
                column: unit_cell_for_rebuild(value)
                for column, value in zip(table.columns, row, strict=True)
                if column != "status"
            }
            for row in table.rows
        ],
    )
    assert rebuilt == frame


def test_read_frame1_table_columns_and_rows_are_tuples_of_tuples():
    frame = _frame1_source_frame()
    table = read_frame1_table(write_frame1((frame,)).data)
    assert isinstance(table.columns, tuple)
    assert isinstance(table.rows, tuple)
    assert all(isinstance(row, tuple) for row in table.rows)


def test_read_frame1_table_refuses_a_wrong_frame_version():
    document = json.dumps({"frame_version": "frame0", "columns": ["stat_var"], "rows": []}).encode(
        "utf-8"
    )
    with pytest.raises(ValueError, match="frame_version"):
        read_frame1_table(document)


def test_read_frame1_table_refuses_a_document_missing_a_required_key():
    document = json.dumps({"frame_version": "frame1", "columns": ["stat_var"]}).encode("utf-8")
    with pytest.raises(ValueError, match="rows"):
        read_frame1_table(document)


def test_read_frame1_table_refuses_a_ragged_row():
    document = json.dumps(
        {"frame_version": "frame1", "columns": ["stat_var", "value"], "rows": [["dcid:X"]]}
    ).encode("utf-8")
    with pytest.raises(ValueError, match="row 0"):
        read_frame1_table(document)


def test_frame1_entity_columns_on_a_two_role_frame_matches_the_sorted_role_keys():
    role_keys = ("recipient", "donor")
    frame = build_frame(
        role_keys=role_keys,
        rows=[
            {
                "stat_var": "dcid:MalariaGrantDisbursements",
                "recipient": "country/TGO",
                "donor": "country/USA",
                "time_period": "2015",
                "value": 1000.0,
            }
        ],
    )
    table = read_frame1_table(write_frame1((frame,)).data)
    assert frame1_entity_columns(table) == tuple(sorted(role_keys))


def test_frame1_entity_columns_refuses_a_table_missing_a_marker_column():
    table = _Table(columns=("recipient", "value"), rows=())
    with pytest.raises(ValueError, match="stat_var"):
        frame1_entity_columns(table)


# --- streaming multi-chunk behaviour -----------------------------------------

_TWO_ROW_CHUNKS = (
    _Table(columns=_GOLDEN_TABLE.columns, rows=(_GOLDEN_TABLE.rows[0],)),
    _Table(columns=_GOLDEN_TABLE.columns, rows=(_GOLDEN_TABLE.rows[1],)),
)


def test_write_csv_over_multiple_chunks_matches_one_whole_chunk():
    whole = write_csv((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    streamed = write_csv(_TWO_ROW_CHUNKS, rendering=_FRIENDLY_RENDERING)
    assert streamed.data == whole.data
    assert streamed.digest == whole.digest


def test_write_excel_over_multiple_chunks_matches_one_whole_chunk():
    whole = write_excel((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    streamed = write_excel(_TWO_ROW_CHUNKS, rendering=_FRIENDLY_RENDERING)
    assert streamed.data == whole.data
    assert streamed.digest == whole.digest


def test_write_parquet_over_multiple_chunks_matches_one_whole_chunk():
    whole = write_parquet((_GOLDEN_TABLE,), rendering=_FRIENDLY_RENDERING)
    streamed = write_parquet(_TWO_ROW_CHUNKS, rendering=_FRIENDLY_RENDERING)
    assert streamed.data == whole.data
    assert streamed.digest == whole.digest


def test_write_csv_preserves_row_order_across_chunks():
    # CSV rows must preserve the order of the three input chunks.
    chunks = (
        _Table(columns=("only",), rows=(("first-in-stream",),)),
        _Table(columns=("only",), rows=(("second",),)),
        _Table(columns=("only",), rows=(("third",),)),
    )
    artifact = write_csv(chunks, rendering=_PLAIN_RENDERING)
    rows = list(csv.reader(io.StringIO(artifact.data.decode("utf-8"))))
    assert rows == [["only"], ["first-in-stream"], ["second"], ["third"]]


def test_write_csv_stringifies_a_mixed_type_any_column_before_sanitizing():
    # The mixed column becomes text. Negative numbers then need formula
    # sanitisation like any other string.
    table = _Table(columns=("mixed",), rows=(("text",), (-5,)))
    artifact = write_csv((table,), rendering=_PLAIN_RENDERING)
    rows = list(csv.reader(io.StringIO(artifact.data.decode("utf-8"))))
    assert rows == [["mixed"], ["text"], ["'-5"]]


def test_write_csv_floats_a_mixed_int_and_float_number_column():
    # A mixed integer/float number column renders integer 1000 as `1000.0`.
    table = _Table(columns=("value",), rows=((1000,), (2500.5,)))
    artifact = write_csv((table,), rendering=_PLAIN_RENDERING)
    rows = list(csv.reader(io.StringIO(artifact.data.decode("utf-8"))))
    assert rows == [["value"], ["1000.0"], ["2500.5"]]


# Groupby and pivot tables have pinned CSV digests, captured from the pre-streaming write_csv,
# so the streaming writer must stay byte-identical.

_GROUPBY_RESULT_TABLE = _Table(
    columns=("stat_var", "recipient", "time_period", "value", "unit"),
    rows=(
        ("dcid:MalariaGrantDisbursements", "country/TGO", "2015", 1000, "USD"),
        ("dcid:MalariaGrantDisbursements", "country/KEN", "2016", 2500.5, "USD"),
    ),
)
"""A groupby SUM aggregate's own output shape: an integer-valued group sum
(`1000`) next to a genuinely fractional one (`2500.5`) in the same `value`
column, widening it to Table Schema's "number" kind."""

_PIVOT_RESULT_TABLE = _Table(
    columns=("recipient", "2015", "2016"),
    rows=(
        ("country/TGO", 1000, "n/a"),
        ("country/KEN", "n/a", 2500.5),
    ),
)
"""A pivot's own output shape: a pivoted-out year column mixing a numeric
observation with a string placeholder for the recipient/year combination
the pivot has no observation for, widening that column to Table Schema's
"any" kind."""


def test_write_csv_groupby_result_matches_the_pre_streaming_golden_digest():
    artifact = write_csv((_GROUPBY_RESULT_TABLE,), rendering=_FRIENDLY_RENDERING)
    assert artifact.digest == (
        "sha256:85fbb81fe9c19c266609895871cffc63b425027ef3cd2fcc9d001a64400b2fe7"
    )
    assert artifact.digest == _recompute_digest(artifact.data)


def test_write_csv_pivot_result_matches_the_pre_streaming_golden_digest():
    artifact = write_csv((_PIVOT_RESULT_TABLE,), rendering=_PLAIN_RENDERING)
    assert artifact.digest == (
        "sha256:0f9595719a74296514da65ed6200e98f7db105df7f5aa9666fe05d26512723c3"
    )
    assert artifact.digest == _recompute_digest(artifact.data)


@pytest.mark.parametrize(
    "call",
    [
        lambda: write_csv((), rendering=_PLAIN_RENDERING),
        lambda: write_excel((), rendering=_PLAIN_RENDERING),
        lambda: write_parquet((), rendering=_PLAIN_RENDERING),
    ],
    ids=["write_csv", "write_excel", "write_parquet"],
)
def test_writers_refuse_an_empty_chunk_stream(call):
    with pytest.raises(MaterialisationUsageError, match="at least one chunk"):
        call()


_MISMATCHED_CHUNKS = (
    _Table(columns=("a", "b"), rows=(("1", "2"),)),
    _Table(columns=("a", "c"), rows=(("3", "4"),)),
)


@pytest.mark.parametrize(
    "call",
    [
        lambda: write_csv(_MISMATCHED_CHUNKS, rendering=_PLAIN_RENDERING),
        lambda: write_excel(_MISMATCHED_CHUNKS, rendering=_PLAIN_RENDERING),
        lambda: write_parquet(_MISMATCHED_CHUNKS, rendering=_PLAIN_RENDERING),
    ],
    ids=["write_csv", "write_excel", "write_parquet"],
)
def test_writers_refuse_a_later_chunk_with_columns_unlike_the_first(call):
    with pytest.raises(MaterialisationUsageError) as exc_info:
        call()
    message = str(exc_info.value)
    assert "chunk 1" in message
    assert "('a', 'c')" in message
    assert "('a', 'b')" in message


def test_write_parquet_coalesces_many_small_chunks_into_row_groups_of_the_target_size(
    monkeypatch,
):
    # Twenty-five rows with a 10-row target form groups of 10, 10, and 5.
    monkeypatch.setattr(materialise_formats, "_PARQUET_ROW_GROUP_ROWS", 10)
    chunks = tuple(_Table(columns=("value",), rows=((f"row-{i}",),)) for i in range(25))
    artifact = write_parquet(chunks, rendering=_PLAIN_RENDERING, size_bound=10_000_000)
    parquet_file = pq.ParquetFile(io.BytesIO(artifact.data))
    metadata = parquet_file.metadata
    assert metadata.num_rows == 25
    assert metadata.num_row_groups == 3
    assert [metadata.row_group(i).num_rows for i in range(3)] == [10, 10, 5]
    # Row order is preserved across row groups.
    assert parquet_file.read().column("value").to_pylist() == [f"row-{i}" for i in range(25)]


def test_write_excel_refuses_over_row_cap_when_the_breach_falls_in_a_later_chunk(
    monkeypatch,
):
    # The header and the first chunk fill the cap, and the second chunk's one row
    # crosses it, so only a count that runs across chunks refuses.
    row_cap = 2
    monkeypatch.setattr(materialise_formats, "EXCEL_MAX_ROWS", row_cap)
    first_chunk = _Table(columns=("only",), rows=tuple((None,) for _ in range(row_cap - 1)))
    second_chunk = _Table(columns=("only",), rows=((None,),))
    with pytest.raises(ExcelLimitError) as exc_info:
        write_excel((first_chunk, second_chunk), rendering=_PLAIN_RENDERING)
    assert str(row_cap + 1) in str(exc_info.value)
