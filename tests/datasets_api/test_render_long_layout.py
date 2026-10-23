"""Tests that `render_and_store` writes long output in the published layout
(`materialise.projection`) while groupby collapse frames keep their grammar columns and
content versions: names beside every code, the split unit, and failed or missing names left
empty with a warning. Every name read of one render shares one deadline and stops after the
first failure. A frame-layout request skips the long layout and reads no names."""

from __future__ import annotations

import csv
import io
import json
import time
import zipfile
from collections import Counter
from collections.abc import Callable, Iterator, Mapping, Sequence

import openpyxl
import pyarrow.parquet as pq
import pytest
from dfo_frame_grammar import Frame
from dfo_frame_grammar.grammar import UnitRequest

from code_index import CodeIndex, CodeRow, CodeSystemDeclaration, IndexStamp
from code_index.shared import get_index
from code_index.types import PartyRow, RoleAxis, StatVarRow, WholeMarker
from datasets_api.manifest_models import CitationSource
from datasets_api.mint_models import GroupbyOptions, MintOptions
from datasets_api.output_options import movable_totals_roles
from datasets_api.render import RenderedDefinition, render_and_store
from dc_graph import GraphDeadlineExceededError, GraphReadError
from materialise.enrichment import EnrichmentLookup
from materialise.frame_build import ParquetFrameChunks, build_frame
from materialise.reshape import Table

from ._artifact_doubles import FakeArtifactStore
from ._graph_doubles import FakeMintGraphReader
from ._render_inputs import PREPARED_AT, SITE

_STAMP = IndexStamp(
    graph_endpoint="https://example.test/core/api",
    exported_at="2026-09-09T00:00:00Z",
    code_counts={},
    statvar_shape_counts={},
    content_hash="deadbeef",
)


def _code(dcid: str, code_system: str, code: str, label: str) -> CodeRow:
    return CodeRow(
        dcid=dcid, code_system=code_system, code=code, label=label, synonyms=(), parent=None
    )


_HEALTH = _code("dcid:CRSPurposeCode122", "crsPurpose", "122", "Basic health care")
_MALARIA = _code("dcid:CRSPurposeCode12250", "crsPurpose", "12250", "Malaria control")
_GRANT = _code("dcid:CRSTypeOfFinanceCode110", "crsTypeOfFinance", "110", "Standard grant")
_LOAN = _code("dcid:CRSTypeOfFinanceCode421", "crsTypeOfFinance", "421", "Standard loan")


def _declaration(code_system: str) -> CodeSystemDeclaration:
    return CodeSystemDeclaration(
        code_system=code_system,
        hierarchical=True,
        disjoint_from_children=True,
        summable=True,
        absent_reading="not_declared",
        plain_name=code_system,
    )


def _statvar(dcid: str, **constraints: str) -> StatVarRow:
    return StatVarRow(
        dcid=dcid,
        population_type="Amount",
        measured_property="amount",
        stat_type=None,
        measurement_qualifier="Nominal",
        measurement_denominator=None,
        observation_properties=("financingProvider", "financingRecipient"),
        constraint_properties=constraints,
    )


_INDEX = CodeIndex(
    stamp=_STAMP,
    declarations=(_declaration("crsPurpose"), _declaration("crsTypeOfFinance")),
    codes=(_HEALTH, _MALARIA, _GRANT, _LOAN),
    statvars=(
        _statvar("SV_A", crsPurpose=_HEALTH.dcid, crsTypeOfFinance=_GRANT.dcid),
        _statvar("SV_B", crsPurpose=_HEALTH.dcid, crsTypeOfFinance=_LOAN.dcid),
        _statvar("SV_C", crsPurpose=_MALARIA.dcid, crsTypeOfFinance=_GRANT.dcid),
        _statvar("SV_D", crsPurpose=_MALARIA.dcid),
    ),
)

_ENTITY_COLUMNS = ("financingProvider", "financingRecipient")
_FRAME_COLUMNS = (
    "stat_var",
    "financingProvider",
    "financingRecipient",
    "time_period",
    "value",
    "unit",
    "measurementMethod",
    "provenance",
    "status",
)
_COLLAPSE_ROWS = (
    ("dcid:SV_A", "country/USA", "country/KEN", "2022", 10.0, "USD", None, "p1", None),
    ("dcid:SV_B", "country/USA", "country/KEN", "2022", 5.0, "USD", None, "p1", None),
    ("dcid:SV_C", "country/USA", "country/KEN", "2022", 7.0, "USD", None, "p1", None),
)
_ROWS = (
    *_COLLAPSE_ROWS,
    ("dcid:SV_D", "country/USA", "country/KEN", "2022", 3.0, "USD", None, "p1", None),
)
_ARCS = {
    "country/USA": {"name": ("United States",)},
    "country/KEN": {"name": ("Kenya",), "countryAlpha3Code": ("KEN",)},
    "p1": {"name": ("Source one",)},
    "WDI_GDP": {"name": ("GDP",)},
}
_GRAPH = FakeMintGraphReader(arcs=_ARCS)
_PARTIES = {
    "country/USA": PartyRow(dcid="country/USA", kind="party", party_class="official"),
    "country/KEN": PartyRow(dcid="country/KEN", kind="party", party_class="official"),
    "DACAfrica": PartyRow(dcid="DACAfrica", kind="total", datasets=("p1",), axes=("destination",)),
    "DACAfricaUnallocated": PartyRow(
        dcid="DACAfricaUnallocated", kind="residual", party_class="residual"
    ),
    "DACDevelopingCountries": PartyRow(
        dcid="DACDevelopingCountries", kind="total", datasets=("p1",), axes=("destination",)
    ),
}
_BY_PURPOSE = GroupbyOptions(by=["financingProvider", "financingRecipient", "crsPurpose"], op="sum")

_SITE_POST = [*_FRAME_COLUMNS, "financingProvider_label"]
"""What the website's `mintColumns` posted once a viewer ticked "Entity names": the output
profile's frame columns, `status` included, then the token."""

_LONG_HEADER = (
    "stat_var",
    "financingProvider",
    "financingProvider_label",
    "financingProvider_kind",
    "financingRecipient",
    "financingRecipient_label",
    "financingRecipient_kind",
    "crsPurpose",
    "crsPurpose_label",
    "crsTypeOfFinance",
    "crsTypeOfFinance_label",
    "time_period",
    "value",
    "unit",
    "unit_label",
    "currency",
    "price_basis",
    "base_year",
    "measurementMethod",
    "measurementMethod_label",
    "provenance",
    "provenance_label",
)
_COLLAPSE_HEADER = (*_LONG_HEADER[:9], *_LONG_HEADER[11:])
"""`_LONG_HEADER` after a groupby by purpose sums over the type of finance."""


def _render(
    options: MintOptions,
    *,
    rows: Sequence[tuple[object, ...]] = _ROWS,
    chunks: Sequence[Sequence[tuple[object, ...]]] | None = None,
    graph: FakeMintGraphReader = _GRAPH,
    index: CodeIndex = _INDEX,
    citations: dict[str, CitationSource] | None = None,
    deadline: float | None = None,
    spec: dict[str, object] | None = None,
    source: str = "tuples",
    parties: dict[str, PartyRow] = _PARTIES,
) -> tuple[RenderedDefinition, bytes]:
    """`chunks` render as tuple tables, as the Arrow frames a single-member mint streams with
    `source="arrow"`, or as the row groups of a group mint's Parquet frame with
    `source="parquet"`."""
    chunks = chunks if chunks is not None else (rows,)
    rows = [row for chunk in chunks for row in chunk]
    store = FakeArtifactStore()
    rendered = render_and_store(
        frame_chunks=_frame_chunks(chunks, source=source),
        options=options,
        spec=spec or {"spec_id": "family"},
        provenance_refs=["p1"],
        citations=citations or {},
        reproducing_call=None,
        content_version="sha256:" + "0" * 64,
        prepared_at=PREPARED_AT,
        site=SITE,
        artifact_store=store,
        entity_columns=_ENTITY_COLUMNS,
        enrichment=EnrichmentLookup(graph, parties=lambda: parties),
        pivot_max_output_cells=5_000_000,
        code_index=index,
        resolved_stat_var_refs=tuple(str(row[0]) for row in rows),
        deadline=deadline,
    )
    return rendered, store.fetch(rendered.artifact.url)


def _frame(rows: Sequence[tuple[object, ...]]) -> Frame:
    return build_frame(
        role_keys=_ENTITY_COLUMNS,
        rows=[
            {column: cell for column, cell in zip(_FRAME_COLUMNS, row, strict=True) if cell}
            for row in rows
        ],
    )


def _frame_chunks(
    chunks: Sequence[Sequence[tuple[object, ...]]], *, source: str
) -> Callable[[], Iterator[Table | Frame]]:
    if source == "tuples":
        return lambda: iter(Table(columns=_FRAME_COLUMNS, rows=tuple(rows)) for rows in chunks)
    frames = [_frame(rows) for rows in chunks]
    if source == "arrow":
        return lambda: iter(frames)
    sink = io.BytesIO()
    with pq.ParquetWriter(sink, frames[0].table.schema) as writer:
        for frame in frames:
            writer.write_table(frame.table)
    return ParquetFrameChunks(sink.getvalue(), columns=_FRAME_COLUMNS)


def _csv(body: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(body.decode("utf-8"))))


def _records(body: bytes) -> list[dict[str, str]]:
    header, *rows = _csv(body)
    return [dict(zip(header, row, strict=True)) for row in rows]


def _header(format_: str, body: bytes) -> tuple[str, ...]:
    if format_ == "csv":
        return tuple(_csv(body)[0])
    if format_ == "parquet":
        return tuple(pq.read_table(io.BytesIO(body)).column_names)
    sheet = openpyxl.load_workbook(io.BytesIO(body), read_only=True).worksheets[0]
    return tuple(next(sheet.iter_rows(max_row=1, values_only=True)))


def _unavailable(rendered: RenderedDefinition) -> list[tuple[object, object]]:
    return [
        (warning.get("column"), warning.get("count"))
        for warning in rendered.warnings
        if warning["code"] == "enrichment_unavailable"
    ]


@pytest.mark.parametrize("format_", ["csv", "xlsx", "parquet"])
def test_every_format_writes_the_long_layout(format_):
    _rendered, body = _render(MintOptions(format=format_, bundle=False))

    assert _header(format_, body) == _LONG_HEADER


@pytest.mark.parametrize(
    ("method", "method_arcs", "method_label"),
    [
        pytest.param(None, {}, "", id="no-method"),
        pytest.param(
            "MOSPIProvisionalEstimates",
            {"name": ("MOSPI Provisional Estimates",)},
            "MOSPI Provisional Estimates",
            id="method-with-a-name",
        ),
        pytest.param(
            "WorldBank_WDI_CSV", {"name": ("WorldBank_WDI_CSV",)}, "", id="method-named-by-its-dcid"
        ),
        pytest.param("GrantEquivalent", {}, "", id="method-without-a-name"),
    ],
)
def test_every_code_gets_its_name_and_the_render_warns_of_nothing(
    method, method_arcs, method_label
):
    rows = [(*row[:6], method, *row[7:]) for row in _ROWS]
    graph = FakeMintGraphReader(arcs={**_ARCS, **({method: method_arcs} if method else {})})

    rendered, body = _render(MintOptions(format="csv", bundle=False), rows=rows, graph=graph)
    first = _records(body)[0]

    assert (first["financingProvider_label"], first["financingRecipient_label"]) == (
        "United States",
        "Kenya",
    )
    assert (first["financingProvider_kind"], first["financingRecipient_kind"]) == (
        "official",
        "country or organisation",
    )
    assert first["provenance_label"] == "Source one"
    # A method shows only a name that says more than its dcid. An unnamed method is not a miss.
    assert (first["measurementMethod"], first["measurementMethod_label"]) == (
        method or "",
        method_label,
    )
    assert _unavailable(rendered) == []
    assert rendered.labels_complete


def test_a_cited_source_takes_its_cited_name_over_the_graph_name():
    _rendered, body = _render(
        MintOptions(format="csv", bundle=False),
        citations={"p1": CitationSource(ref="p1", label="OECD Creditor Reporting System")},
    )

    assert {record["provenance_label"] for record in _records(body)} == {
        "OECD Creditor Reporting System"
    }


_NOMINAL_SHAPE_SPEC: dict[str, object] = {
    "spec_id": "gdp",
    "identity": {"shape": {"kind": "statvar_schema", "measurement_qualifier": "Nominal"}},
}
"""A described identity whose shape states current prices, over StatVars the code index does
not hold."""


@pytest.mark.parametrize(
    ("unit", "stat_var", "spec", "expected"),
    [
        pytest.param(
            "USD", None, None, ("USD", "Current US dollars", "USD", "current", ""), id="native"
        ),
        pytest.param(
            "undata/UNIT_MEASURE-CR_USD",
            None,
            None,
            ("undata/UNIT_MEASURE-CR_USD", "Current US dollars", "USD", "current", ""),
            id="declared-un-dollars",
        ),
        pytest.param(
            "undata/UNIT_MEASURE-CR_USD_K",
            None,
            None,
            ("undata/UNIT_MEASURE-CR_USD_K", "Current US dollars", "USD", "current", ""),
            id="declared-un-thousand-dollars",
        ),
        pytest.param(
            "ChainedUSDollarBaseYr2017",
            None,
            None,
            ("ChainedUSDollarBaseYr2017", "Constant 2017 US dollars", "USD", "constant", "2017"),
            id="declared-chained-2017",
        ),
        pytest.param(
            "USDollar",
            "dcid:WDI_GDP",
            _NOMINAL_SHAPE_SPEC,
            ("USDollar", "Current US dollars", "USD", "current", ""),
            id="native-outside-the-index",
        ),
        pytest.param(
            "INR",
            "dcid:WDI_GDP",
            _NOMINAL_SHAPE_SPEC,
            ("INR", "Current Indian rupees", "INR", "current", ""),
            id="native-graph-currency-outside-the-index",
        ),
        pytest.param(
            UnitRequest(currency="USD", prices="constant", base_year=2023),
            None,
            None,
            ("CurrencyUSD", "Constant 2023 US dollars", "USD", "constant", "2023"),
            id="converted-constant",
        ),
        pytest.param(
            UnitRequest(currency="KES", prices="current"),
            None,
            None,
            ("CurrencyKES", "Current Kenyan shillings", "KES", "current", ""),
            id="converted-current",
        ),
        pytest.param(
            "Percent", None, None, ("Percent", "Percent", "", "", ""), id="unnamed-non-currency"
        ),
    ],
)
def test_the_unit_splits_into_node_label_currency_basis_and_year(unit, stat_var, spec, expected):
    rows = [(stat_var or row[0], *row[1:5], unit, *row[6:]) for row in _ROWS]

    rendered, body = _render(MintOptions(format="csv", bundle=False), rows=rows, spec=spec)

    columns = ("unit", "unit_label", "currency", "price_basis", "base_year")
    assert {tuple(record[c] for c in columns) for record in _records(body)} == {expected}
    assert _unavailable(rendered) == []


def test_a_statvar_without_a_value_on_an_axis_gets_an_empty_code_and_the_absent_reading():
    _rendered, body = _render(MintOptions(format="csv", bundle=False))
    by_stat_var = {record["stat_var"]: record for record in _records(body)}

    assert by_stat_var["dcid:SV_A"]["crsTypeOfFinance"] == _GRANT.dcid
    assert by_stat_var["dcid:SV_A"]["crsTypeOfFinance_label"] == "Standard grant"
    assert by_stat_var["dcid:SV_D"]["crsTypeOfFinance"] == ""
    assert by_stat_var["dcid:SV_D"]["crsTypeOfFinance_label"] == "not_declared"


@pytest.mark.parametrize(
    ("groupby", "rows", "columns", "header"),
    [
        pytest.param(None, _ROWS, _SITE_POST, _LONG_HEADER, id="frame-columns-and-token"),
        pytest.param(
            _BY_PURPOSE, _COLLAPSE_ROWS, _SITE_POST, _COLLAPSE_HEADER, id="groupby-collapse"
        ),
        pytest.param(None, _ROWS, ["financingProvider_label"], _LONG_HEADER, id="label-token"),
        pytest.param(
            None,
            _ROWS,
            ["financingRecipient_iso3", "financingProvider_label"],
            (*_LONG_HEADER[:7], "financingRecipient_iso3", *_LONG_HEADER[7:]),
            id="iso3-and-label-tokens",
        ),
        # The kind of each party role is a default column; a token names it once.
        pytest.param(None, _ROWS, ["party_kind"], _LONG_HEADER, id="bare-kind-token"),
        pytest.param(
            None, _ROWS, ["financingRecipient_party_kind"], _LONG_HEADER, id="qualified-kind-token"
        ),
        # A selection that omits `financingProvider` but names its label keeps both.
        pytest.param(
            None,
            _ROWS,
            [
                "stat_var",
                "financingRecipient",
                "time_period",
                "value",
                "unit",
                "provenance",
                "financingProvider_label",
            ],
            (*_LONG_HEADER[:18], *_LONG_HEADER[20:]),
            id="selection-naming-a-label-keeps-its-entity",
        ),
    ],
)
def test_a_posted_label_token_selects_nothing_beyond_its_entity(groupby, rows, columns, header):
    rendered, body = _render(
        MintOptions(format="csv", bundle=False, groupby=groupby, columns=columns), rows=rows
    )

    assert tuple(_csv(body)[0]) == header
    assert _unavailable(rendered) == []


def test_a_groupby_collapse_keeps_its_frame_and_content_version():
    rendered, body = _render(
        MintOptions(format="csv", bundle=False, groupby=_BY_PURPOSE), rows=_COLLAPSE_ROWS
    )

    assert rendered.collapse is not None
    # The collapse frame and its identity are the grammar's, untouched by the written columns.
    assert rendered.collapse.content_version == "cv_03c9050c06fd518c9cfb8c7178a06033"
    assert rendered.collapse.frame.columns == (
        "stat_var",
        "financingProvider",
        "financingRecipient",
        "stat_var_crsPurpose",
        "time_period",
        "value",
        "unit",
        "measurementMethod",
        "provenance",
        "status",
    )
    header, *rows = _csv(body)
    assert tuple(header) == _COLLAPSE_HEADER
    assert [row[7:10] for row in rows] == [
        [_HEALTH.dcid, "Basic health care", "2022"],
        [_MALARIA.dcid, "Malaria control", "2022"],
    ]
    # The summed family's own native unit keeps the price basis its members share.
    assert {row[header.index("price_basis")] for row in rows} == {"current"}


class _FailingGraph(FakeMintGraphReader):
    """A graph whose batched arc read fails after its first `answers` reads, recording the
    deadline each read carried. With a `clock`, a failing read holds until its deadline, as
    one does in a graph outage."""

    def __init__(
        self, error: GraphReadError, *, clock: _Clock | None = None, answers: int = 0
    ) -> None:
        super().__init__(arcs=_ARCS)
        self._error = error
        self._clock = clock
        self._answers = answers
        self.deadlines: list[float | None] = []

    def node_arcs_many(self, dcids, *, properties, individual_probe_budget=None, deadline=None):
        self.deadlines.append(deadline)
        if len(self.deadlines) <= self._answers:
            return super().node_arcs_many(
                dcids,
                properties=properties,
                individual_probe_budget=individual_probe_budget,
                deadline=deadline,
            )
        if self._clock is not None and deadline is not None:
            self._clock.now = max(self._clock.now, deadline)
        raise self._error


class _SlowGraph(FakeMintGraphReader):
    """A graph whose batched arc read answers after `read_seconds` on `clock`, recording the
    deadline each read carried."""

    def __init__(self, *, clock: _Clock, read_seconds: float) -> None:
        super().__init__(arcs={**_ARCS, "country/GHA": {"name": ("Ghana",)}})
        self._clock = clock
        self._read_seconds = read_seconds
        self.deadlines: list[float | None] = []

    def node_arcs_many(self, dcids, *, properties, individual_probe_budget=None, deadline=None):
        self.deadlines.append(deadline)
        self._clock.now += self._read_seconds
        return super().node_arcs_many(
            dcids,
            properties=properties,
            individual_probe_budget=individual_probe_budget,
            deadline=deadline,
        )


class _Clock:
    """A monotonic clock moved only by the graph doubles."""

    def __init__(self) -> None:
        self.now = 1_000.0


@pytest.fixture
def clock() -> _Clock:
    return _Clock()


@pytest.mark.parametrize(
    "error",
    [
        pytest.param(GraphReadError("graph unavailable"), id="read-error"),
        pytest.param(GraphDeadlineExceededError("deadline passed"), id="deadline"),
    ],
)
def test_a_failed_name_read_leaves_name_cells_empty_warns_and_completes(error):
    graph = _FailingGraph(error)

    rendered, body = _render(
        MintOptions(format="csv", bundle=False), graph=graph, deadline=time.monotonic()
    )

    records = _records(body)
    assert {record["financingProvider_label"] for record in records} == {""}
    assert {record["provenance_label"] for record in records} == {""}
    # The unit's label and the parties' index kinds are local, so they survive the failed read.
    assert {record["unit_label"] for record in records} == {"Current US dollars"}
    assert {record["financingRecipient_kind"] for record in records} == {"country or organisation"}
    assert _unavailable(rendered) == [
        ("financingProvider_label", 4),
        ("financingRecipient_label", 4),
        ("provenance_label", 4),
    ]
    # One read, bounded by the metadata deadline, and the artifact is linked as incomplete.
    assert len(graph.deadlines) == 1
    assert graph.deadlines[0] is not None
    assert not rendered.labels_complete


def test_a_label_two_codes_of_a_system_share_leaves_the_code_empty_and_warns():
    twin = _code("dcid:CRSPurposeCode99999", "crsPurpose", "99999", "Basic health care")
    index = CodeIndex(
        stamp=_STAMP,
        declarations=_INDEX.declarations.values(),
        codes=(*_INDEX.codes, twin),
        statvars=_INDEX.statvars,
    )

    rendered, body = _render(MintOptions(format="csv", bundle=False), index=index)

    by_stat_var = {record["stat_var"]: record for record in _records(body)}
    assert by_stat_var["dcid:SV_A"]["crsPurpose"] == ""
    assert by_stat_var["dcid:SV_A"]["crsPurpose_label"] == "Basic health care"
    assert by_stat_var["dcid:SV_C"]["crsPurpose"] == _MALARIA.dcid
    assert _unavailable(rendered) == [("crsPurpose", 2)]


def test_no_fixture_code_system_gives_two_codes_one_label():
    """An axis code column reads its code back from the decoded label, so a label two codes
    of one declared system share leaves those rows without a code."""
    index = get_index()
    labels = Counter(
        (row.code_system, row.label) for row in index.codes if row.code_system in index.declarations
    )

    assert [key for key, count in labels.items() if count > 1] == []


def test_a_decode_gap_warning_names_the_label_column():
    gap_index_statvar = "dcid:SV_GAP"
    rows = (
        (gap_index_statvar, "country/USA", "country/KEN", "2022", 1.0, "USD", None, "p1", None),
    )
    index = CodeIndex(
        stamp=_STAMP,
        declarations=(_declaration("crsPurpose"),),
        codes=(_HEALTH,),
        statvars=(_statvar("SV_GAP", crsPurpose="dcid:CRSPurposeCode99999"),),
    )

    rendered, _body = _render(MintOptions(format="csv", bundle=False), rows=rows, index=index)

    assert _unavailable(rendered) == [("crsPurpose_label", 1)]


def test_every_chunk_of_a_render_reads_names_against_one_deadline(clock):
    """A later chunk's new names are read against the deadline the render started with, so a
    slow first read leaves later reads less time."""
    graph = _SlowGraph(clock=clock, read_seconds=5.0)
    later = [(*row[:2], "country/GHA", *row[3:]) for row in _ROWS]

    rendered, _body = _render(
        MintOptions(format="csv", bundle=False),
        chunks=(_ROWS, later),
        graph=graph,
        deadline=clock.now + 1.0,
    )

    # The render's own deadline, for both chunks' reads.
    assert graph.deadlines == [1_001.0, 1_001.0]
    assert rendered.labels_complete


@pytest.mark.parametrize("bundle", [False, True], ids=["bare", "bundle"])
def test_a_render_stops_reading_names_after_the_first_failed_read(clock, bundle):
    """A graph outage costs a render one read's wait, whatever its chunk count: no later
    chunk with new names, and no bundle codes table, reads again."""
    graph = _FailingGraph(GraphDeadlineExceededError("deadline passed"), clock=clock)
    started = clock.now
    later = [(*row[:2], "country/GHA", *row[3:]) for row in _ROWS]

    rendered, _body = _render(
        MintOptions(format="csv", bundle=bundle),
        chunks=(_ROWS, later),
        graph=graph,
        deadline=started + 1.0,
    )

    assert graph.deadlines == [started + 1.0]
    assert clock.now == started + 1.0
    assert not rendered.labels_complete


_FRAME_HEADER = (*_FRAME_COLUMNS, "stat_var_crsPurpose", "stat_var_crsTypeOfFinance")
_AXIS_LABELS = (
    ("Basic health care", "Standard grant"),
    ("Basic health care", "Standard loan"),
    ("Malaria control", "Standard grant"),
    ("Malaria control", "not_declared"),
)
_CONSTANT_2023_USD = UnitRequest(currency="USD", prices="constant", base_year=2023)


@pytest.mark.parametrize(
    ("extra", "rows", "header", "expected"),
    [
        pytest.param(
            {},
            _ROWS,
            _FRAME_HEADER,
            [(*row, *labels) for row, labels in zip(_ROWS, _AXIS_LABELS, strict=True)],
            id="decoded-axes",
        ),
        pytest.param(
            {
                "groupby": {
                    "by": ["financingProvider", "financingRecipient", "crsPurpose"],
                    "op": "sum",
                }
            },
            _COLLAPSE_ROWS,
            (*_FRAME_COLUMNS[:3], "stat_var_crsPurpose", *_FRAME_COLUMNS[3:]),
            [
                (
                    "dcid:family",
                    "country/USA",
                    "country/KEN",
                    "Basic health care",
                    "2022",
                    15.0,
                    "USD",
                    None,
                    "p1",
                    None,
                ),
                (
                    "dcid:family",
                    "country/USA",
                    "country/KEN",
                    "Malaria control",
                    "2022",
                    7.0,
                    "USD",
                    None,
                    "p1",
                    None,
                ),
            ],
            id="summed",
        ),
        pytest.param(
            {},
            [(*row[:5], _CONSTANT_2023_USD, *row[6:]) for row in _ROWS],
            _FRAME_HEADER,
            [
                (
                    *row[:5],
                    '{"base_year":2023,"currency":"USD","prices":"constant"}',
                    *row[6:],
                    *labels,
                )
                for row, labels in zip(_ROWS, _AXIS_LABELS, strict=True)
            ],
            id="converted-unit",
        ),
        pytest.param(
            {"columns": ["financingRecipient_iso3"]},
            _ROWS,
            (*_FRAME_HEADER, "financingRecipient_iso3"),
            [(*row, *labels, "KEN") for row, labels in zip(_ROWS, _AXIS_LABELS, strict=True)],
            id="enrichment-column",
        ),
    ],
)
def test_a_frame_layout_file_holds_the_frame_columns_and_cells(extra, rows, header, expected):
    """The columns and cells a bare parquet file held before the long layout, which the
    pre-layout writer produced byte for byte for the same frames. Its rendering is the request
    options as sent, with no layout version stamped."""
    options = {"format": "parquet", "bundle": False, "layout": "frame", **extra}

    rendered, body = _render(MintOptions.model_validate(options), rows=rows)

    table = pq.read_table(io.BytesIO(body))
    assert tuple(table.column_names) == header
    assert [tuple(row.values()) for row in table.to_pylist()] == expected
    assert rendered.artifact.rendering == options


def test_a_frame_layout_render_reads_no_names():
    graph = _FailingGraph(GraphReadError("graph unavailable"))

    rendered, _body = _render(
        MintOptions(format="parquet", bundle=False, layout="frame"), graph=graph
    )

    assert graph.deadlines == []
    assert rendered.labels_complete
    assert _unavailable(rendered) == []


_RECIPIENT_ROWS = (
    ("dcid:SV_A", "country/USA", "country/KEN", "2021", 10.0, "USD", None, "p1", None),
    ("dcid:SV_A", "country/USA", "DACAfricaUnallocated", "2021", 2.0, "USD", None, "p1", None),
    ("dcid:SV_A", "country/USA", "DACAfrica", "2021", 12.0, "USD", None, "p1", None),
    ("dcid:SV_A", "country/USA", "DACDevelopingCountries", "2022", 40.0, "USD", None, "p1", None),
)
_PARTY_ROWS = ["country/KEN", "DACAfricaUnallocated"]
_TOTAL_ROWS = ["DACAfrica", "DACDevelopingCountries"]


def _spec(identity: dict[str, object]) -> dict[str, object]:
    return {"spec_id": "family", "identity": {"shape": {"kind": "statvar_schema"}, **identity}}


_OPEN_RECIPIENTS = _spec({"entities": {"origin": ["country/USA"]}, "open_roles": ["destination"]})


def _recipients(records: list[dict[str, str]]) -> list[str]:
    return [record["financingRecipient"] for record in records]


def _sheet_rows(body: bytes, sheet: str) -> list[tuple[object, ...]]:
    workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True)
    return [tuple(row) for row in workbook[sheet].iter_rows(values_only=True)]


@pytest.mark.parametrize(
    ("spec", "rows", "kept"),
    [
        pytest.param(_OPEN_RECIPIENTS, _RECIPIENT_ROWS, _PARTY_ROWS, id="open-role-drops-totals"),
        pytest.param(
            _spec({"entities": {"recipient": ["country/KEN", "DACAfrica"]}}),
            _RECIPIENT_ROWS[::2],
            ["country/KEN"],
            id="several-bound-drop-totals",
        ),
        pytest.param(
            _spec(
                {
                    "entities": {
                        "origin": ["country/USA"],
                        "destination": ["DACDevelopingCountries"],
                    }
                }
            ),
            _RECIPIENT_ROWS[3:],
            ["DACDevelopingCountries"],
            id="one-bound-total-keeps-every-row",
        ),
        pytest.param(None, _RECIPIENT_ROWS, [*_PARTY_ROWS, *_TOTAL_ROWS], id="no-identity"),
    ],
)
def test_a_bare_csv_leaves_out_the_totals_of_an_open_or_several_bound_role(spec, rows, kept):
    """ "ODA from the UK to developing countries, total" binds the recipient to one total, so
    its rows are the data and stay."""
    _rendered, body = _render(MintOptions(format="csv", bundle=False), rows=rows, spec=spec)

    assert _recipients(_records(body)) == kept


def _totals_published(
    *, datasets: tuple[str, ...], axes: tuple[RoleAxis, ...], whole: tuple[WholeMarker, ...] = ()
) -> dict[str, PartyRow]:
    return {
        **_PARTIES,
        **{
            dcid: PartyRow(dcid=dcid, kind="total", datasets=datasets, axes=axes, whole=whole)
            for dcid in _TOTAL_ROWS
        },
    }


_WHOLE_TOTALS = _totals_published(
    datasets=("p1",),
    axes=("destination",),
    whole=(WholeMarker(dataset="p1", covers=("official", "residual")),),
)


@pytest.mark.parametrize(
    ("parties", "named", "kept"),
    [
        pytest.param(_PARTIES, ["financingRecipient"], _PARTY_ROWS, id="read-provenance-totals"),
        pytest.param(
            _WHOLE_TOTALS, ["financingRecipient"], _PARTY_ROWS, id="read-provenance-whole-totals"
        ),
        pytest.param(
            _totals_published(datasets=("p1",), axes=("origin",)),
            [],
            [*_PARTY_ROWS, *_TOTAL_ROWS],
            id="totals-only-as-donors",
        ),
        pytest.param(
            _totals_published(datasets=("p2",), axes=("destination",)),
            [],
            [*_PARTY_ROWS, *_TOTAL_ROWS],
            id="totals-of-another-provenance",
        ),
    ],
)
def test_a_bare_csv_drops_totals_only_on_the_roles_the_options_document_names(parties, named, kept):
    """The render splits by `movable_totals_roles`, the rule behind `long_layout_totals`, so a
    row the index declares a total stays when its role is not advertised."""
    _rendered, body = _render(
        MintOptions(format="csv", bundle=False),
        rows=_RECIPIENT_ROWS,
        spec=_OPEN_RECIPIENTS,
        parties=parties,
    )

    identity = _OPEN_RECIPIENTS["identity"]
    assert (
        list(movable_totals_roles(identity, _ENTITY_COLUMNS, parties=parties, datasets=["p1"]))
        == named
    )
    assert _recipients(_records(body)) == kept


@pytest.mark.parametrize("bundle", [False, True])
@pytest.mark.parametrize("parties", [_PARTIES, _WHOLE_TOTALS], ids=["totals", "whole-totals"])
def test_an_xlsx_holds_the_totals_on_a_second_sheet_and_years_as_numbers(bundle, parties):
    _rendered, body = _render(
        MintOptions(format="xlsx", bundle=bundle),
        rows=_RECIPIENT_ROWS,
        spec=_OPEN_RECIPIENTS,
        parties=parties,
    )
    if bundle:
        archive = zipfile.ZipFile(io.BytesIO(body))
        body = archive.read("data/p1.xlsx")
        (resource,) = [
            entry
            for entry in json.loads(archive.read("datapackage.json"))["resources"]
            if entry["name"] != "codes"
        ]
        assert "on the sheet totals" in resource["description"]
    recipient = _LONG_HEADER.index("financingRecipient")
    year = _LONG_HEADER.index("time_period")

    data = _sheet_rows(body, "data")
    totals = _sheet_rows(body, "totals")

    assert data[0] == totals[0] == _LONG_HEADER
    assert [row[recipient] for row in data[1:]] == _PARTY_ROWS
    assert [row[recipient] for row in totals[1:]] == _TOTAL_ROWS
    assert [row[year] for row in totals[1:]] == [2021, 2022]


def test_a_bundled_csv_writes_the_totals_as_a_second_data_file():
    _rendered, body = _render(
        MintOptions(format="csv", bundle=True), rows=_RECIPIENT_ROWS, spec=_OPEN_RECIPIENTS
    )
    archive = zipfile.ZipFile(io.BytesIO(body))
    resources = {
        entry["path"]: entry
        for entry in json.loads(archive.read("datapackage.json"))["resources"]
        if entry["name"] != "codes"
    }

    assert _recipients(_records(archive.read("data/p1.csv"))) == _PARTY_ROWS
    assert _recipients(_records(archive.read("data/p1-totals.csv"))) == _TOTAL_ROWS
    assert list(resources) == ["data/p1.csv", "data/p1-totals.csv"]
    assert "data/p1-totals.csv" in resources["data/p1.csv"]["description"]
    codes = {row["code"] for row in _records(archive.read("data/codes.csv"))}
    assert codes >= {"country/KEN", "DACAfrica", "DACDevelopingCountries"}


@pytest.mark.parametrize("format_", ["csv", "xlsx"])
def test_an_open_role_holding_only_totals_leaves_the_data_file_empty_and_says_so(format_):
    """The totals type the columns of the resource that holds them: the CSV totals file, or
    the one XLSX file, whose data sheet is empty."""
    _rendered, body = _render(
        MintOptions(format=format_, bundle=True), rows=_RECIPIENT_ROWS[2:], spec=_OPEN_RECIPIENTS
    )
    archive = zipfile.ZipFile(io.BytesIO(body))

    if format_ == "csv":
        assert _csv(archive.read("data/p1.csv")) == [list(_LONG_HEADER)]
        assert _recipients(_records(archive.read("data/p1-totals.csv"))) == _TOTAL_ROWS
        assert "No rows: every observation is of a total" in archive.read("README.md").decode()
        typed_path = "data/p1-totals.csv"
    else:
        workbook = archive.read("data/p1.xlsx")
        assert _sheet_rows(workbook, "data") == [_LONG_HEADER]
        totals = _sheet_rows(workbook, "totals")
        assert [row[_LONG_HEADER.index("financingRecipient")] for row in totals[1:]] == _TOTAL_ROWS
        typed_path = "data/p1.xlsx"
    resource = next(
        entry
        for entry in json.loads(archive.read("datapackage.json"))["resources"]
        if entry["path"] == typed_path
    )
    types = {field["name"]: field["type"] for field in resource["schema"]["fields"]}
    assert (types["time_period"], types["value"]) == ("integer", "number")


@pytest.mark.parametrize(
    ("options", "header_at"),
    [
        pytest.param(MintOptions(format="parquet", bundle=False), _LONG_HEADER, id="parquet"),
        pytest.param(
            MintOptions(
                format="csv",
                bundle=False,
                pivot={"index": "financingRecipient", "columns": "time_period"},
            ),
            None,
            id="pivot",
        ),
    ],
)
def test_parquet_and_a_pivot_keep_every_row(options, header_at):
    _rendered, body = _render(options, rows=_RECIPIENT_ROWS, spec=_OPEN_RECIPIENTS)

    if options.format == "parquet":
        table = pq.read_table(io.BytesIO(body))
        assert table.column("financingRecipient").to_pylist() == [*_PARTY_ROWS, *_TOTAL_ROWS]
        assert str(table.schema.field("time_period").type) == "int64"
    else:
        assert [record["financingRecipient"] for record in _records(body)] == [
            *_PARTY_ROWS,
            *_TOTAL_ROWS,
        ]


def test_totals_spread_over_chunks_keep_the_first_written_chunk_typed():
    """A chunk of totals alone empties the main table's first chunk. The writer still types
    its columns from a chunk with rows, so a bundle's schema reads numbers."""
    _rendered, body = _render(
        MintOptions(format="csv", bundle=True),
        chunks=(_RECIPIENT_ROWS[2:], _RECIPIENT_ROWS[:2]),
        spec=_OPEN_RECIPIENTS,
    )
    archive = zipfile.ZipFile(io.BytesIO(body))
    resource = next(
        entry
        for entry in json.loads(archive.read("datapackage.json"))["resources"]
        if entry["path"] == "data/p1.csv"
    )
    types = {field["name"]: field["type"] for field in resource["schema"]["fields"]}

    assert _recipients(_records(archive.read("data/p1.csv"))) == _PARTY_ROWS
    assert (types["value"], types["time_period"]) == ("number", "integer")


@pytest.mark.parametrize(
    ("chunks", "answers", "columns"),
    [
        pytest.param((_RECIPIENT_ROWS,), 0, None, id="first-read-fails"),
        pytest.param((_RECIPIENT_ROWS[:2], _RECIPIENT_ROWS[2:]), 1, None, id="later-read-fails"),
        pytest.param((_RECIPIENT_ROWS,), 0, ["party_kind"], id="bare-kind-token"),
        pytest.param(
            (_RECIPIENT_ROWS,), 0, ["financingRecipient_party_kind"], id="qualified-kind-token"
        ),
    ],
)
def test_a_failed_name_read_still_sets_the_index_totals_apart(chunks, answers, columns):
    """The code index gives a listed party its kind with no graph read, so a graph outage from
    the first read or partway through the render moves every listed total, in every chunk. A
    token naming a party role's kind gives the same default column."""
    graph = _FailingGraph(GraphReadError("graph unavailable"), answers=answers)

    rendered, body = _render(
        MintOptions(format="csv", bundle=True, columns=columns),
        chunks=chunks,
        graph=graph,
        spec=_OPEN_RECIPIENTS,
        deadline=time.monotonic() + 60.0,
    )
    archive = zipfile.ZipFile(io.BytesIO(body))
    data = _records(archive.read("data/p1.csv"))
    totals = _records(archive.read("data/p1-totals.csv"))
    resources = {
        entry["path"]: entry for entry in json.loads(archive.read("datapackage.json"))["resources"]
    }

    assert len(graph.deadlines) == answers + 1
    assert not rendered.labels_complete
    assert _recipients(data) == _PARTY_ROWS
    assert _recipients(totals) == _TOTAL_ROWS
    assert [record["financingRecipient_kind"] for record in data] == [
        "country or organisation",
        "unallocated",
    ]
    assert [record["financingRecipient_kind"] for record in totals] == ["total", "total"]
    assert "data/p1-totals.csv" in resources["data/p1.csv"]["description"]
    assert "data/p1-totals.csv" in archive.read("README.md").decode()


_UNLISTED_RECIPIENT_ROW = (
    "dcid:SV_A",
    "country/USA",
    "country/GHA",
    "2022",
    4.0,
    "USD",
    None,
    "p1",
    None,
)
_PIVOT_BY_YEAR = {"index": "financingRecipient", "columns": "time_period"}


@pytest.mark.parametrize(
    ("options", "rows", "kinds", "reads"),
    [
        pytest.param(
            MintOptions(format="csv", bundle=False, columns=["party_kind"], pivot=_PIVOT_BY_YEAR),
            _RECIPIENT_ROWS,
            None,
            0,
            id="pivot",
        ),
        pytest.param(
            MintOptions(format="parquet", bundle=False, layout="frame", columns=["party_kind"]),
            _RECIPIENT_ROWS,
            ["party", "residual", "total", "total"],
            0,
            id="frame-bare",
        ),
        pytest.param(
            MintOptions(
                format="parquet",
                bundle=False,
                layout="frame",
                columns=["financingRecipient_party_kind"],
            ),
            (*_RECIPIENT_ROWS, _UNLISTED_RECIPIENT_ROW),
            ["party", "residual", "total", "total", ""],
            0,
            id="frame-qualified-unlisted-party",
        ),
    ],
)
def test_a_party_kind_token_outside_the_long_layout_reads_the_index_alone(
    options, rows, kinds, reads
):
    """The reshape enrichment reads a party's kind from the code index, as the long layout
    does, and never reads the graph for it. A party the index does not list has an empty
    cell with a warning, and the render completes."""
    graph = _FailingGraph(GraphReadError("graph unavailable"))

    rendered, body = _render(
        options, rows=rows, graph=graph, spec=_OPEN_RECIPIENTS, deadline=time.monotonic() + 60.0
    )

    assert len(graph.deadlines) == reads
    assert rendered.labels_complete == (reads == 0)
    if kinds is not None:
        column = pq.read_table(io.BytesIO(body)).column("financingRecipient_party_kind")
        assert column.to_pylist() == kinds
    assert _unavailable(rendered) == (
        [("financingRecipient_party_kind", 1)] if "" in (kinds or ()) else []
    )


_MIXED_PERIOD_ROWS = (
    ("dcid:SV_A", "country/USA", "country/KEN", "2021", 10.0, "USD", None, "p1", None),
    ("dcid:SV_A", "country/USA", "country/KEN", "2022", 11.0, "USD", None, "p1", None),
    ("dcid:SV_A", "country/USA", "country/KEN", "2019/2021", 30.0, "USD", None, "p1", None),
)


@pytest.mark.parametrize(
    ("chunks", "typed"),
    [
        pytest.param((_MIXED_PERIOD_ROWS[:2],), True, id="year-points-are-integers"),
        pytest.param((_MIXED_PERIOD_ROWS,), False, id="one-chunk"),
        pytest.param(
            (_MIXED_PERIOD_ROWS[:2], _MIXED_PERIOD_ROWS[2:]), False, id="span-in-a-later-chunk"
        ),
        pytest.param((_MIXED_PERIOD_ROWS[::-1],), False, id="span-first"),
    ],
)
@pytest.mark.parametrize("source", ["tuples", "arrow", "parquet"])
@pytest.mark.parametrize("format_", ["csv", "parquet"])
def test_a_frame_mixing_year_points_and_spans_writes_every_period_as_text(
    format_, source, chunks, typed, monkeypatch
):
    periods = sorted(str(row[3]) for chunk in chunks for row in chunk)
    written = sorted(int(period) for period in periods) if typed else periods
    passes: list[int] = []
    read_frames = ParquetFrameChunks.__call__

    def counted(self: ParquetFrameChunks) -> Iterator[Frame]:
        passes.append(1)
        return read_frames(self)

    monkeypatch.setattr(ParquetFrameChunks, "__call__", counted)

    _rendered, body = _render(
        MintOptions(format=format_, bundle=format_ == "csv"), chunks=chunks, source=source
    )
    if source == "parquet":
        # The grain is read from the `time_period` column alone, so only the write decodes rows.
        assert passes == [1]

    if format_ == "parquet":
        column = pq.read_table(io.BytesIO(body)).column("time_period")
        assert str(column.type) == ("int64" if typed else "string")
        assert sorted(column.to_pylist()) == written
        return
    archive = zipfile.ZipFile(io.BytesIO(body))
    (resource,) = [
        entry
        for entry in json.loads(archive.read("datapackage.json"))["resources"]
        if entry["name"] != "codes"
    ]
    types = {field["name"]: field["type"] for field in resource["schema"]["fields"]}
    assert types["time_period"] == ("integer" if typed else "string")
    assert sorted(r["time_period"] for r in _records(archive.read(resource["path"]))) == periods


_DAC2A = "provenance/OECDDAC2A"
_DAC2B = "provenance/OECDDAC2B"
_DAC_PARTIES = {
    **_PARTIES,
    "DACAllRecipients": PartyRow(
        dcid="DACAllRecipients",
        kind="total",
        datasets=(_DAC2A, _DAC2B),
        axes=("destination",),
        party_class="official",
        whole=(WholeMarker(dataset=_DAC2A, covers=("official", "residual"), years=(2020, 2021)),),
    ),
    "DACOfficialDonors": PartyRow(
        dcid="DACOfficialDonors",
        kind="total",
        datasets=(_DAC2A, _DAC2B),
        axes=("origin",),
        party_class="official",
        whole=(
            WholeMarker(dataset=_DAC2A, covers=("official",)),
            WholeMarker(dataset=_DAC2B, covers=("official",)),
        ),
    ),
    "GatesFoundation": PartyRow(dcid="GatesFoundation", kind="party", party_class="private"),
}


def _recipient_row(recipient: str, period: str, provenance: str) -> tuple[object, ...]:
    return ("dcid:SV_A", "country/USA", recipient, period, 1.0, "USD", None, provenance, None)


_WINDOWED_ROWS = (
    _recipient_row("country/KEN", "2021", _DAC2A),
    _recipient_row("DACAllRecipients", "2019", _DAC2A),
    _recipient_row("DACAllRecipients", "2021", _DAC2A),
    _recipient_row("DACAllRecipients", "2021-06", _DAC2A),
    _recipient_row("DACAllRecipients", "2021/2022", _DAC2A),
    _recipient_row("DACAllRecipients", "2022", _DAC2A),
    _recipient_row("DACAllRecipients", "2021", _DAC2B),
    _recipient_row("DACAllRecipients", "2020", _DAC2B),
)
_WINDOWED_WHOLE = [None, False, True, True, False, False, False, False]
"""Whether each total of `_WINDOWED_ROWS` is whole there, in order (`None` for the party). The
total is whole in DAC2A inside its marker's years (2020 to 2021), and not in DAC2A outside
them, across years, or in any DAC2B row, where it carries no marker."""

_WINDOWED_KINDS = [
    "country or organisation",
    *("whole total" if whole else "total" for whole in _WINDOWED_WHOLE[1:]),
]
_WINDOWED_CELLS = [
    ("party", "false"),
    ("total", "false"),
    ("total", "true"),
    ("total", "true"),
    ("total", ""),
    ("total", "false"),
    ("total", "false"),
    ("total", "false"),
]
"""The frame cells of `_WINDOWED_ROWS`, which follow `_WINDOWED_WHOLE` except on the DAC2A row
across 2021 and 2022: that row names no one year for the marker's window, so its `party_whole`
is empty."""


def _kinds_by_row(rows: Sequence[Mapping[str, object]]) -> list[object]:
    return [row["financingRecipient_kind"] for row in rows]


def _cells_by_row(rows: Sequence[Mapping[str, object]]) -> list[object]:
    return [
        (row["financingRecipient_party_kind"], row["financingRecipient_party_whole"])
        for row in rows
    ]


@pytest.mark.parametrize(
    ("options", "reader", "expected"),
    [
        pytest.param(
            MintOptions(format="csv", bundle=False),
            lambda body: _kinds_by_row(_records(body)),
            _WINDOWED_KINDS,
            id="long-layout",
        ),
        pytest.param(
            MintOptions(
                format="parquet",
                bundle=False,
                layout="frame",
                columns=["financingRecipient_party_kind", "financingRecipient_party_whole"],
            ),
            lambda body: _cells_by_row(pq.read_table(io.BytesIO(body)).to_pylist()),
            _WINDOWED_CELLS,
            id="reshape-enrichment-keeping-time-period",
        ),
    ],
)
def test_a_total_is_whole_only_in_the_dataset_and_years_of_its_marker(options, reader, expected):
    """One frame holds DAC2A and DAC2B rows of the same total in several periods: whether it is
    whole follows each row's dataset and year, in the long layout and in the reshape enrichment
    alike."""
    _rendered, body = _render(options, rows=_WINDOWED_ROWS, spec=None, parties=_DAC_PARTIES)

    assert reader(body) == expected


def test_a_whole_cell_a_pivot_indexes_by_is_empty_for_a_windowed_whole_total():
    """A pivot groups rows by its index, so a cell that changed with the year would split one
    recipient across index rows. The pivot reads the cell for no dataset and year, and a whole
    total's cell is empty there."""
    options = MintOptions(
        format="csv",
        bundle=False,
        groupby=GroupbyOptions(by=["financingProvider", "financingRecipient"], op="sum"),
        columns=[
            "stat_var",
            "financingProvider",
            "financingRecipient",
            "financingRecipient_party_whole",
            "time_period",
            "value",
            "unit",
            "provenance",
        ],
        pivot={"index": "financingRecipient_party_whole", "columns": "time_period"},
    )
    one_row_per_year = (_WINDOWED_ROWS[1], _WINDOWED_ROWS[2], _WINDOWED_ROWS[5])

    _rendered, body = _render(options, rows=one_row_per_year, spec=None, parties=_DAC_PARTIES)

    header, *rows = _csv(body)
    assert [row[header.index("financingRecipient_party_whole")] for row in rows] == [""]


def test_a_dac2a_ethiopia_2021_download_reads_official_private_and_whole_total_providers():
    rows = tuple(
        ("dcid:SV_A", provider, "country/ETH", "2021", 1.0, "USD", None, _DAC2A, None)
        for provider in ("country/USA", "GatesFoundation", "DACOfficialDonors")
    )

    _rendered, body = _render(
        MintOptions(format="csv", bundle=False), rows=rows, spec=None, parties=_DAC_PARTIES
    )

    assert [record["financingProvider_kind"] for record in _records(body)] == [
        "official",
        "private",
        "whole total",
    ]
