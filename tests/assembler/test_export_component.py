"""Checks that a component export serves exactly the frame rows behind the values the card
shows, renders them through the Datasets writers, and bundles their citation, codes and column
meanings as a valid Frictionless package and CFF citation, named and dated by the revision."""

from __future__ import annotations

import csv
import dataclasses
import io
import json
import zipfile
from collections.abc import Mapping
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import frictionless
import jsonschema
import openpyxl
import pyarrow.parquet as pq
import pytest
import yaml
from dfo_frame_grammar import Frame, canonical_frame_bytes

from assembler.errors import RevisionNotFound
from assembler.fakes import FakeRevisionStore
from assembler.frame_view import compute_structural_view
from assembler.pipeline import write_component
from assembler.pipeline.export_component import (
    ComponentExportMismatch,
    ComponentHoldsNoData,
    ComponentPredatesRowIdentities,
    export_component_data,
)
from assembler.ports import FrameRef, ResolvedFrame
from assembler.testing import FIXTURE_RETRIEVED_ON, build_resolved_frame
from materialise.bundle_naming import card_slug
from tests.presentation_store.test_frame_store import (
    _FakeGraphReader,
    _store,
    _thin_spec_fixture,
)

_FRAME_REF = FrameRef(
    spec_id="spec_9f2e41a7c03b58d2a4c1f6b2e8d35790",
    content_version="cv_f4a81c22d9075e3ba61c2ee98f83b405",
)
_OPS_FRAME_REF = FrameRef(spec_id="ops_" + "c" * 32, content_version="ocv_" + "d" * 32)
_STAT_VAR = "dcid:MalariaGrantDisbursements"
_SITE_URL = "https://site.example"
# 23:30 on 4 October at UTC-5 is 5 October in UTC, the date file names carry.
_CREATED_AT = datetime(2026, 10, 4, 23, 30, tzinfo=timezone(timedelta(hours=-5)))
_CFF_SCHEMA = json.loads(
    (
        Path(__file__).parents[1] / "materialise/fixtures/bundle-citation-cff-1.2.0-schema.json"
    ).read_text()
)
_OECD_ATTRIBUTION = "OECD (2026), Creditor Reporting System, OECD Data Explorer."
_FORMULA_LABEL = '=HYPERLINK("http://example.test","Kenya")'


def _row(*, recipient, period, value, provenance="src_oecd_dac", stat_var=_STAT_VAR):
    return {
        "stat_var": stat_var,
        "recipient": recipient,
        "time_period": period,
        "value": value,
        "unit": None,
        "measurementMethod": None,
        "provenance": provenance,
    }


def _frame(frame_ref: FrameRef = _FRAME_REF) -> ResolvedFrame:
    """Three recipients over 2019-2023, GHA from a second source, ETH missing 2021, and rows of
    another stat var and outside every window that no card shows."""
    rows = [
        _row(recipient=recipient, period=str(year), value=float(base + year - 2019))
        for recipient, base in (("country/ETH", 100), ("country/KEN", 300))
        for year in range(2019, 2024)
        if not (recipient == "country/ETH" and year == 2021)
    ]
    rows += [
        _row(recipient="country/GHA", period=str(year), value=200.0 + year, provenance="src_imf")
        for year in range(2019, 2024)
    ]
    rows += [_row(recipient="country/ETH", period="2020", value=9.0, stat_var="dcid:Other")]
    return build_resolved_frame(
        frame_ref=frame_ref,
        role_keys=("recipient",),
        rows=rows,
        labels={
            "country/ETH": "Ethiopia",
            "country/GHA": "Ghana",
            "country/KEN": _FORMULA_LABEL,
            _STAT_VAR: "Malaria grant disbursements",
        },
        manifest={
            "schema_version": 1,
            "citation": {
                "text": "exemplar",
                "sources": [
                    {
                        "ref": "src_oecd_dac",
                        "label": "OECD DAC",
                        "license": "CC0",
                        "license_type": "CreativeCommons0",
                        "license_url": "https://creativecommons.org/publicdomain/zero/1.0/",
                        "url": "https://oecd.example/dac",
                        "source": {"ref": "OECD", "label": "OECD"},
                        "documentation": {
                            "publisher": None,
                            "attribution": _OECD_ATTRIBUTION,
                            "last_update": "2026-10-02",
                            "currencies": ["USD"],
                            "transformation_log": [
                                {"transformation": "Aggregate totals", "description": "Sums."}
                            ],
                            "version_history": [
                                {
                                    "date": "2026-09-01",
                                    "title": "First release",
                                    "changes": [{"kind": "added", "note": "The source."}],
                                }
                            ],
                        },
                    },
                    {
                        "ref": "src_imf",
                        "label": "IMF",
                        "license": "IMF terms",
                        "license_url": "https://imf.example/terms",
                    },
                ],
            },
        },
    )


@pytest.fixture
def revision_store() -> FakeRevisionStore:
    return FakeRevisionStore(clock=lambda: _CREATED_AT)


def _frame_ref_body(frame_ref: FrameRef) -> dict[str, str]:
    return {"spec_id": frame_ref.spec_id, "content_version": frame_ref.content_version}


_INTENTS: Mapping[str, Mapping[str, object]] = {
    "line_chart": {
        "type": "line_chart",
        "select": {
            "stat_var": _STAT_VAR,
            "recipient": ["country/ETH", "country/KEN"],
            "provenance": "src_oecd_dac",
        },
        "window": {"start": "2020", "end": "2022"},
    },
    "bar_chart": {
        "type": "bar_chart",
        "select": {
            "stat_var": _STAT_VAR,
            "recipient": ["country/ETH", "country/KEN"],
            "provenance": "src_oecd_dac",
        },
        "window": {"start": "2019", "end": "2023"},
    },
    "table": {
        "type": "table",
        "select": {
            "stat_var": _STAT_VAR,
            "recipient": ["country/ETH", "country/GHA"],
            "provenance": ["src_oecd_dac", "src_imf"],
        },
        "window": {"start": "2020", "end": "2022"},
    },
    "stat_card": {
        "type": "stat_card",
        "select": {"stat_var": _STAT_VAR, "recipient": "country/KEN", "provenance": "src_oecd_dac"},
        "change": {"basis": ["2019", "2022"], "derivation": "percent_change"},
    },
}


def _create(
    *, member, frame_store, revision_store, release_registry, frame_ref: FrameRef = _FRAME_REF
):
    frame_store.put(resolved_frame=_frame(frame_ref))
    intent = {**_INTENTS[member], "frame_ref": _frame_ref_body(frame_ref)}
    return write_component(
        intent=intent,
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )


def _export(created, *, frame_store, revision_store, format="csv", bundle=False):
    return export_component_data(
        component_id=created.component_id,
        revision=created.revision,
        format=format,
        bundle=bundle,
        revision_store=revision_store,
        frame_store=frame_store,
        site_url=_SITE_URL,
    )


def _unzip(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return {name: zf.read(name) for name in zf.namelist()}


def _validated(tmp_path: Path, files: Mapping[str, bytes]) -> dict:
    """The bundle's descriptor, after frictionless validates the unzipped package and its
    CITATION.cff validates against the CFF 1.2.0 schema."""
    for name, data in files.items():
        path = tmp_path / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
    report = frictionless.validate(str(tmp_path / "datapackage.json"))
    assert report.valid, report.flatten(["type", "note"])
    jsonschema.validate(instance=yaml.safe_load(files["CITATION.cff"]), schema=_CFF_SCHEMA)
    return json.loads(files["datapackage.json"])


def _field_titles(descriptor: dict) -> dict[str, str]:
    (resource, _codes) = descriptor["resources"]
    return {field["name"]: field["title"] for field in resource["schema"]["fields"]}


def _csv_rows(data: bytes) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(data.decode("utf-8"))))


def _neutralised(text: str) -> str:
    return "'" + text if text.startswith("=") else text


def _shown(component: Mapping[str, object]) -> list[tuple[str, str]]:
    """The `(series label, period)` of every value the card shows, in card order, with labels
    as the CSV writer neutralises them."""
    if component["type"] == "table":
        return [
            (_neutralised(" / ".join(row["dimensions"].values())), row["period"])
            for row in component["rows"]
            if "value" in row
        ]
    if component["type"] == "stat_card":
        periods = dict.fromkeys([component["period"], *component["change"]["basis"]])
        return [(_neutralised(component["label"]), period) for period in periods]
    return [
        (_neutralised(series["label"]), point["period"])
        for series in component["series"]
        for point in series["points"]
        if "value" in point
    ]


@pytest.mark.parametrize("member", ["line_chart", "bar_chart", "table", "stat_card"])
def test_export_holds_one_frame_row_per_value_the_card_shows(
    member, frame_store, revision_store, release_registry
):
    created = _create(
        member=member,
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )
    component = created.resolved_payload["component"]

    exported = _export(created, frame_store=frame_store, revision_store=revision_store)
    rows = _csv_rows(exported.data)

    assert [(row["series"], row["time_period"]) for row in rows] == _shown(component)
    assert {row["stat_var"] for row in rows} == {_STAT_VAR}
    assert exported.media_type == "text/csv"
    assert exported.file_name == (f"dfo_{card_slug(component['title']['text'])}_2026-10-05_r1.csv")


def test_entity_bar_exports_one_row_per_bar_at_the_shared_period(
    frame_store, revision_store, release_registry
):
    created = _create(
        member="bar_chart",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    rows = _csv_rows(_export(created, frame_store=frame_store, revision_store=revision_store).data)

    assert [(row["recipient"], row["time_period"], row["value"]) for row in rows] == [
        ("country/ETH", "2023", "104.0"),
        ("country/KEN", "2023", "304.0"),
    ]


@pytest.mark.parametrize(
    ("member", "recipient", "shown"),
    [
        ("bar_chart", ["country/ETH", "country/KEN"], 2),
        ("table", ["country/ETH", "country/KEN"], 2),
        ("stat_card", "country/KEN", 1),
    ],
)
def test_a_span_frame_exports_the_span_as_each_row_s_period(
    member, recipient, shown, frame_store, revision_store, release_registry
):
    frame = build_resolved_frame(
        frame_ref=_FRAME_REF,
        role_keys=("recipient",),
        rows=[
            _row(recipient="country/ETH", period="2018/2022", value=500.0),
            _row(recipient="country/KEN", period="2018/2022", value=300.0),
        ],
    )
    frame_store.put(resolved_frame=frame)
    created = write_component(
        intent={
            "type": member,
            "frame_ref": _frame_ref_body(_FRAME_REF),
            "select": {"stat_var": _STAT_VAR, "recipient": recipient},
        },
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    rows = _csv_rows(_export(created, frame_store=frame_store, revision_store=revision_store).data)

    assert [row["time_period"] for row in rows] == ["2018/2022"] * shown


def test_export_columns_follow_the_frame_with_labels_beside_their_refs(
    frame_store, revision_store, release_registry
):
    created = _create(
        member="stat_card",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    data = _export(created, frame_store=frame_store, revision_store=revision_store).data
    header = next(csv.reader(io.StringIO(data.decode("utf-8"))))

    assert header == [
        "series",
        "stat_var",
        "stat_var_label",
        "recipient",
        "recipient_label",
        "time_period",
        "value",
        "unit",
        "measurementMethod",
        "provenance",
        "provenance_label",
        "status",
    ]


def test_a_label_starting_with_a_formula_trigger_is_neutralised_in_csv_and_literal_in_xlsx(
    frame_store, revision_store, release_registry
):
    created = _create(
        member="stat_card",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    csv_rows = _csv_rows(
        _export(created, frame_store=frame_store, revision_store=revision_store).data
    )
    xlsx = _export(created, frame_store=frame_store, revision_store=revision_store, format="xlsx")
    sheet = openpyxl.load_workbook(io.BytesIO(xlsx.data)).active
    header = [cell.value for cell in sheet[1]]
    label_cell = sheet.cell(row=2, column=header.index("recipient_label") + 1)

    assert {row["recipient_label"] for row in csv_rows} == {"'" + _FORMULA_LABEL}
    assert {row["series"] for row in csv_rows} == {"'" + _FORMULA_LABEL}
    assert label_cell.data_type == "s"
    assert label_cell.value == "'" + _FORMULA_LABEL


def test_parquet_export_carries_the_same_rows(frame_store, revision_store, release_registry):
    created = _create(
        member="line_chart",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    exported = _export(
        created, frame_store=frame_store, revision_store=revision_store, format="parquet"
    )
    table = pq.read_table(io.BytesIO(exported.data))

    assert table.column("time_period").to_pylist() == [2020, 2022, 2020, 2021, 2022]
    assert table.schema.field("time_period").type == "int64"
    assert table.schema.field("value").type == "double"


def test_bundle_validates_and_documents_its_sources_columns_and_codes(
    tmp_path, frame_store, revision_store, release_registry
):
    created = _create(
        member="table",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )
    exported = _export(created, frame_store=frame_store, revision_store=revision_store, bundle=True)
    files = _unzip(exported.data)
    descriptor = _validated(tmp_path, files)
    citation = yaml.safe_load(files["CITATION.cff"])
    readme = files["README.md"].decode("utf-8")
    sources_md = files["SOURCES.md"].decode("utf-8")
    codes = {row["code"]: (row["label"], row["kind"]) for row in _csv_rows(files["data/codes.csv"])}
    resource, _codes = descriptor["resources"]
    route = f"{_SITE_URL}/api/presentation/components/{created.component_id}/revisions/1/data"

    assert exported.media_type == "application/zip"
    assert exported.file_name == "dfo_stat-var-recipient-provenance_2026-10-05_csv_r1.zip"
    assert sorted(files) == [
        "CITATION.cff",
        "README.md",
        "SOURCES.md",
        "data/codes.csv",
        "data/stat-var-recipient-provenance.csv",
        "datapackage.json",
    ]
    assert len(_csv_rows(files["data/stat-var-recipient-provenance.csv"])) == len(
        _shown(created.resolved_payload["component"])
    )
    assert descriptor["created"] == "2026-10-05T04:30:00Z"
    assert citation["date-released"] == "2026-10-05"
    assert descriptor["id"] == route
    assert descriptor["dfo"]["reproducing_call"]["url"] == route
    assert resource["licenses"] == [
        {
            "name": "CC0-1.0",
            "path": "https://creativecommons.org/publicdomain/zero/1.0/",
            "title": "CC0",
        },
        {"path": "https://imf.example/terms", "title": "IMF terms"},
    ]
    assert [reference["title"] for reference in citation["references"]] == ["OECD DAC", "IMF"]
    assert citation["references"][0]["notes"] == _OECD_ATTRIBUTION
    assert "Aggregate totals" in sources_md and "First release" in sources_md
    assert codes["country/ETH"] == ("Ethiopia", "entity")
    assert codes[_STAT_VAR] == ("Malaria grant disbursements", "variable")
    assert codes["src_oecd_dac"] == ("OECD DAC", "source")
    assert sorted(key["fields"][0] for key in resource["schema"]["foreignKeys"]) == [
        "provenance",
        "recipient",
        "stat_var",
    ]
    assert _field_titles(descriptor) == {
        "series": "Series",
        "stat_var": "Variable",
        "stat_var_label": "Variable name",
        "recipient": "Recipient",
        "recipient_label": "Recipient name",
        "time_period": "Period",
        "value": "Value",
        "unit": "Unit",
        "measurementMethod": "Measurement method",
        "provenance": "Source dataset",
        "provenance_label": "Source dataset name",
        "status": "Status",
    }
    assert descriptor["dfo"]["component"] == {
        "component_id": created.component_id,
        "revision": 1,
        "frame_ref": _frame_ref_body(_FRAME_REF),
        "select": _INTENTS["table"]["select"],
        "window": _INTENTS["table"]["window"],
    }
    assert _FRAME_REF.spec_id in readme and _FRAME_REF.content_version in readme


def test_a_card_over_a_thin_root_frame_exports_the_licence_the_live_fill_reads(
    frame_store, revision_store, release_registry
):
    """The stored citation has null licence fields. The frame store fills them from the graph
    on load, so the bundle's descriptor carries a `licenses` entry the stored citation lacks."""
    snapshots, artifacts, frame_ref, thin = _thin_spec_fixture()
    graph = _FakeGraphReader(
        arcs={
            "src_oecd_dac": {
                "typeOf": ("Provenance",),
                "name": ("OECD Creditor Reporting System",),
                "license": ("https://example.org/terms",),
            }
        }
    )
    resolved = _store(snapshots=snapshots, artifacts=artifacts, graph=graph).load(
        frame_ref=frame_ref
    )
    frame_store.put(resolved_frame=resolved)
    intent = {
        **_INTENTS["line_chart"],
        "select": {
            "stat_var": _STAT_VAR,
            "recipient": ["country/ETH"],
            "provenance": "src_oecd_dac",
        },
        "frame_ref": _frame_ref_body(frame_ref),
    }
    created = write_component(
        intent=intent,
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    exported = _export(created, frame_store=frame_store, revision_store=revision_store, bundle=True)

    descriptor = json.loads(_unzip(exported.data)["datapackage.json"])
    assert thin["sources"][0]["license_url"] is None
    assert descriptor["resources"][0]["licenses"] == [
        {"path": "https://example.org/terms", "title": "CC BY 4.0"}
    ]


@pytest.mark.parametrize("format", ["csv", "xlsx", "parquet"])
@pytest.mark.parametrize("bundle", [True, False])
def test_two_exports_of_one_revision_are_byte_identical_under_one_name(
    format, bundle, frame_store, revision_store, release_registry
):
    created = _create(
        member="line_chart",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    first, second = (
        _export(
            created,
            frame_store=frame_store,
            revision_store=revision_store,
            format=format,
            bundle=bundle,
        )
        for _ in range(2)
    )

    assert first == second


def test_a_revision_read_without_its_creation_time_refuses(
    frame_store, revision_store, release_registry
):
    created = _create(
        member="line_chart",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )
    records = revision_store._revisions[created.component_id]
    records[-1] = dataclasses.replace(records[-1], created_at=None)

    with pytest.raises(RuntimeError):
        _export(created, frame_store=frame_store, revision_store=revision_store, bundle=True)


def test_an_ops_frame_exports(frame_store, revision_store, release_registry):
    created = _create(
        member="line_chart",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
        frame_ref=_OPS_FRAME_REF,
    )

    exported = _export(created, frame_store=frame_store, revision_store=revision_store, bundle=True)
    files = _unzip(exported.data)
    descriptor = json.loads(files["datapackage.json"])
    data_path = descriptor["resources"][0]["path"]

    assert len(_csv_rows(files[data_path])) == 5
    assert descriptor["dfo"]["component"]["frame_ref"] == _frame_ref_body(_OPS_FRAME_REF)


def test_a_frame_value_that_differs_from_the_card_refuses(
    frame_store, revision_store, release_registry
):
    created = _create(
        member="line_chart",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )
    drifted = _frame()
    columns = drifted.frame["columns"]
    rows = [
        [*row[: columns.index("value")], -1.0, *row[columns.index("value") + 1 :]]
        if row[columns.index("recipient")] == "country/ETH"
        and row[columns.index("time_period")] == "2020"
        and row[columns.index("stat_var")] == _STAT_VAR
        else list(row)
        for row in drifted.frame["rows"]
    ]
    frame_store.put(
        resolved_frame=dataclasses.replace(drifted, frame={**drifted.frame, "rows": rows})
    )

    with pytest.raises(ComponentExportMismatch):
        _export(created, frame_store=frame_store, revision_store=revision_store)


def test_a_frame_missing_a_shown_row_refuses(frame_store, revision_store, release_registry):
    created = _create(
        member="stat_card",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )
    frame = _frame()
    columns = frame.frame["columns"]
    kept = [
        row
        for row in frame.frame["rows"]
        if not (
            row[columns.index("recipient")] == "country/KEN"
            and row[columns.index("time_period")] == "2019"
        )
    ]
    frame_store.put(resolved_frame=dataclasses.replace(frame, frame={**frame.frame, "rows": kept}))

    with pytest.raises(ComponentExportMismatch):
        _export(created, frame_store=frame_store, revision_store=revision_store)


def test_a_note_holds_no_data(frame_store, revision_store, release_registry):
    created = write_component(
        intent={"type": "note_narration", "text": "Disbursements rose."},
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    with pytest.raises(ComponentHoldsNoData):
        _export(created, frame_store=frame_store, revision_store=revision_store)


def test_a_revision_stored_without_series_coordinates_holds_no_data(
    frame_store, revision_store, release_registry
):
    created = _create(
        member="line_chart",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )
    records = revision_store._revisions[created.component_id]
    records[-1] = dataclasses.replace(records[-1], series_coordinates=None)

    with pytest.raises(ComponentHoldsNoData):
        _export(created, frame_store=frame_store, revision_store=revision_store)


def test_a_missing_revision_refuses(frame_store, revision_store, release_registry):
    created = _create(
        member="line_chart",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    with pytest.raises(RevisionNotFound):
        export_component_data(
            component_id=created.component_id,
            revision=2,
            format="csv",
            bundle=False,
            revision_store=revision_store,
            frame_store=frame_store,
            site_url=_SITE_URL,
        )


def _strip_table_series_ids(revision_store, created, *, positions) -> None:
    """Remove `series_id` from the stored table rows at `positions`, as a revision stored before
    rows carried it holds them."""
    records = revision_store._revisions[created.component_id]
    payload = records[-1].resolved_payload
    rows = [dict(row) for row in payload["component"]["rows"]]
    for position in positions(rows):
        del rows[position]["series_id"]
    stored = {**payload, "component": {**payload["component"], "rows": rows}}
    records[-1] = dataclasses.replace(records[-1], resolved_payload=stored)


@pytest.mark.parametrize(
    ("positions", "refusal"),
    [
        pytest.param(
            lambda rows: range(len(rows)), ComponentPredatesRowIdentities, id="no-row-has-one"
        ),
        pytest.param(lambda rows: range(1), ComponentExportMismatch, id="one-row-lacks-one"),
    ],
)
def test_a_table_revision_without_row_series_ids_refuses(
    positions, refusal, frame_store, revision_store, release_registry
):
    """A table stored before its rows carried `series_id` holds no traceable rows, and one with
    only some rows carrying it disagrees with itself."""
    created = _create(
        member="table",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )
    _strip_table_series_ids(revision_store, created, positions=positions)

    with pytest.raises(refusal):
        _export(created, frame_store=frame_store, revision_store=revision_store)


@pytest.mark.parametrize("format", ["csv", "xlsx", "parquet"])
def test_generated_columns_take_a_free_name_beside_frame_columns_holding_theirs(
    format, frame_store, revision_store, release_registry
):
    """Roles named `series` and `stat_var_label` keep their names and meanings, and the
    generated series and label columns move to the first free suffix and keep theirs."""
    rows = [
        {
            "stat_var": _STAT_VAR,
            "series": "country/KEN",
            "stat_var_label": "sector/health",
            "time_period": str(year),
            "value": float(year),
            "unit": None,
            "measurementMethod": None,
            "provenance": "src_oecd_dac",
        }
        for year in (2019, 2022)
    ]
    frame_store.put(
        resolved_frame=build_resolved_frame(
            frame_ref=_FRAME_REF, role_keys=("series", "stat_var_label"), rows=rows
        )
    )
    created = write_component(
        intent={
            "type": "stat_card",
            "frame_ref": _frame_ref_body(_FRAME_REF),
            "select": {
                "stat_var": _STAT_VAR,
                "series": "country/KEN",
                "stat_var_label": "sector/health",
                "provenance": "src_oecd_dac",
            },
            "change": {"basis": ["2019", "2022"], "derivation": "percent_change"},
        },
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    data = _export(
        created, frame_store=frame_store, revision_store=revision_store, format=format
    ).data
    if format == "csv":
        header = next(csv.reader(io.StringIO(data.decode("utf-8"))))
    elif format == "xlsx":
        header = [cell.value for cell in openpyxl.load_workbook(io.BytesIO(data)).active[1]]
    else:
        header = pq.read_table(io.BytesIO(data)).column_names

    assert header == [
        "series_2",
        "stat_var",
        "stat_var_label_2",
        "series",
        "series_label",
        "stat_var_label",
        "stat_var_label_label",
        "time_period",
        "value",
        "unit",
        "measurementMethod",
        "provenance",
        "provenance_label",
        "status",
    ]
    bundled = _unzip(
        _export(
            created,
            frame_store=frame_store,
            revision_store=revision_store,
            format=format,
            bundle=True,
        ).data
    )
    titles = _field_titles(json.loads(bundled["datapackage.json"]))
    assert {name: titles[name] for name in header[:7]} == {
        "series_2": "Series",
        "stat_var": "Variable",
        "stat_var_label_2": "Variable name",
        "series": "Entity",
        "series_label": "Entity name",
        "stat_var_label": "Entity",
        "stat_var_label_label": "Entity name",
    }


def test_a_role_a_spreadsheet_escapes_keeps_its_label_column_title(
    tmp_path, frame_store, revision_store, release_registry
):
    """Roles named `-donor` and `-donor_label` are written as `'-donor` and `'-donor_label`. The
    generated column naming `-donor` takes the free name `-donor_label_2`, written as
    `'-donor_label_2`, and the bundle still titles it as that role's name column."""
    rows = [
        {
            "stat_var": _STAT_VAR,
            "-donor": "country/KEN",
            "-donor_label": "country/TGO",
            "time_period": str(year),
            "value": float(year),
            "unit": None,
            "measurementMethod": None,
            "provenance": "src_oecd_dac",
        }
        for year in (2019, 2022)
    ]
    frame_store.put(
        resolved_frame=build_resolved_frame(
            frame_ref=_FRAME_REF, role_keys=("-donor", "-donor_label"), rows=rows
        )
    )
    created = write_component(
        intent={
            "type": "stat_card",
            "frame_ref": _frame_ref_body(_FRAME_REF),
            "select": {
                "stat_var": _STAT_VAR,
                "-donor": "country/KEN",
                "-donor_label": "country/TGO",
                "provenance": "src_oecd_dac",
            },
            "change": {"basis": ["2019", "2022"], "derivation": "percent_change"},
        },
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    bundled = _unzip(
        _export(created, frame_store=frame_store, revision_store=revision_store, bundle=True).data
    )
    titles = _field_titles(_validated(tmp_path, bundled))
    assert {name: titles[name] for name in ("'-donor", "'-donor_label_2")} == {
        "'-donor": "Entity",
        "'-donor_label_2": "Entity name",
    }


_TARGET_FRAME_REF = FrameRef(spec_id=_FRAME_REF.spec_id, content_version="cv_" + "b" * 32)
_MOVED_BY = 1000.0


def _moved_frame() -> ResolvedFrame:
    """`_frame` under `_TARGET_FRAME_REF`, every value `_MOVED_BY` higher."""
    base = _frame(_TARGET_FRAME_REF)
    value = base.frame["columns"].index("value")
    rows = [[*row[:value], row[value] + _MOVED_BY, *row[value + 1 :]] for row in base.frame["rows"]]
    return dataclasses.replace(base, frame={**base.frame, "rows": rows})


@pytest.mark.parametrize("member", ["line_chart", "bar_chart", "table", "stat_card"])
def test_export_on_a_target_frame_holds_that_frame_s_rows_for_the_rendered_card(
    member, frame_store, revision_store, release_registry
):
    created = _create(
        member=member,
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )
    frame_store.put(resolved_frame=_moved_frame())
    stored = _csv_rows(
        _export(created, frame_store=frame_store, revision_store=revision_store).data
    )

    exported = export_component_data(
        component_id=created.component_id,
        revision=created.revision,
        format="csv",
        bundle=False,
        revision_store=revision_store,
        frame_store=frame_store,
        site_url=_SITE_URL,
        target_frame=_TARGET_FRAME_REF,
        release_registry=release_registry,
    )
    rows = _csv_rows(exported.data)

    assert [(row["series"], row["time_period"]) for row in rows] == _shown(
        created.resolved_payload["component"]
    )
    assert [float(row["value"]) for row in rows] == [
        float(row["value"]) + _MOVED_BY for row in stored
    ]
    assert len(revision_store._revisions[created.component_id]) == 1
    assert (_TARGET_FRAME_REF, created.component_id) not in frame_store.pins


def test_a_bundle_on_a_target_frame_is_dated_by_that_frame_and_names_no_reproducing_call(
    tmp_path, frame_store, revision_store, release_registry
):
    """The revision was created on 2026-10-05 and the target frame retrieved on 2026-10-06."""
    created = _create(
        member="table",
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )
    frame_store.put(
        resolved_frame=dataclasses.replace(_moved_frame(), retrieved_on=date(2026, 10, 6))
    )

    exported = export_component_data(
        component_id=created.component_id,
        revision=created.revision,
        format="csv",
        bundle=True,
        revision_store=revision_store,
        frame_store=frame_store,
        site_url=_SITE_URL,
        target_frame=_TARGET_FRAME_REF,
        release_registry=release_registry,
    )
    files = _unzip(exported.data)
    descriptor = _validated(tmp_path, files)

    saved = f"{_SITE_URL}/api/presentation/components/{created.component_id}/revisions/1/data"
    assert exported.file_name == "dfo_stat-var-recipient-provenance_2026-10-06_csv_r1.zip"
    assert descriptor["dfo"]["prepared_at"] == "2026-10-06T00:00:00Z"
    assert descriptor["dfo"]["component"]["frame_ref"] == _frame_ref_body(_TARGET_FRAME_REF)
    assert "reproducing_call" not in descriptor["dfo"]
    assert descriptor["dfo"]["links"]["version"] == saved
    assert "/render/data" not in files["README.md"].decode("utf-8")


_USD_PER_PERSON = {"numerator": "USDollar", "denominator": {"stat_var": "dcid:Count_Person"}}


_USD_PER_PERSON_TEXT = '{"denominator":{"stat_var":"dcid:Count_Person"},"numerator":"USDollar"}'


def _put_composed_unit_frame(frame_store, *, labels: Mapping[str, str]) -> None:
    frame = Frame.build(
        role_keys=("recipient",),
        rows=[
            {
                **_row(recipient="country/KEN", period=str(year), value=float(year)),
                "unit": _USD_PER_PERSON,
            }
            for year in (2019, 2022)
        ],
    )
    # The stored frame1 document holds a composed unit as a JSON object.
    document = json.loads(canonical_frame_bytes(frame))
    frame_store.put(
        resolved_frame=ResolvedFrame(
            frame_ref=_OPS_FRAME_REF,
            frame=document,
            structural_view=compute_structural_view(frame=document, continuation=None),
            manifest={"schema_version": 1},
            labels=labels,
            retrieved_on=FIXTURE_RETRIEVED_ON,
        )
    )


def test_a_table_export_words_a_composed_unit_in_series_and_keeps_its_json_in_unit(
    frame_store, revision_store, release_registry
):
    _put_composed_unit_frame(
        frame_store, labels={"USDollar": "US Dollar", "dcid:Count_Person": "Total population"}
    )
    created = write_component(
        intent={
            "type": "table",
            "frame_ref": _frame_ref_body(_OPS_FRAME_REF),
            "select": {"recipient": "country/KEN", "unit": _USD_PER_PERSON_TEXT},
        },
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    rows = _csv_rows(_export(created, frame_store=frame_store, revision_store=revision_store).data)

    assert [row["series"] for row in rows] == ["country/KEN / US Dollar per Total population"] * 2
    assert [row["unit"] for row in rows] == [_USD_PER_PERSON_TEXT] * 2


@pytest.mark.parametrize("format", ["csv", "xlsx", "parquet"])
def test_a_composed_unit_exports_as_its_canonical_json(
    format, frame_store, revision_store, release_registry
):
    _put_composed_unit_frame(frame_store, labels={})
    created = write_component(
        intent={
            "type": "stat_card",
            "frame_ref": _frame_ref_body(_OPS_FRAME_REF),
            "select": {"stat_var": _STAT_VAR, "recipient": "country/KEN"},
            "change": {"basis": ["2019", "2022"], "derivation": "percent_change"},
        },
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    data = _export(
        created, frame_store=frame_store, revision_store=revision_store, format=format
    ).data
    if format == "csv":
        units = [row["unit"] for row in _csv_rows(data)]
    elif format == "xlsx":
        sheet = openpyxl.load_workbook(io.BytesIO(data)).active
        column = [cell.value for cell in sheet[1]].index("unit")
        units = [row[column] for row in sheet.iter_rows(min_row=2, values_only=True)]
    else:
        units = pq.read_table(io.BytesIO(data)).column("unit").to_pylist()

    assert units == [_USD_PER_PERSON_TEXT] * 2


def _put_continued_frame(frame_store) -> None:
    """ETH reported for 2019-2021 under a null method and continued by its source's estimates
    for 2022-2023 (`trace.continuations`), beside GHA reported under `Reported`."""
    rows = [
        {
            **_row(recipient="country/ETH", period=str(year), value=float(year)),
            "measurementMethod": "Estimated" if year >= 2022 else None,
        }
        for year in range(2019, 2024)
    ]
    rows += [
        {
            **_row(recipient="country/GHA", period=str(year), value=float(year) + 0.5),
            "measurementMethod": "Reported",
        }
        for year in range(2019, 2024)
    ]
    frame_store.put(
        resolved_frame=build_resolved_frame(
            frame_ref=_FRAME_REF,
            role_keys=("recipient",),
            rows=rows,
            labels={"country/ETH": "Ethiopia", "country/GHA": "Ghana"},
            manifest={
                "schema_version": 1,
                "trace": {
                    "continuations": [
                        {
                            "axis": "measurementMethod",
                            "method": "Estimated",
                            "periods": ["2022", "2023"],
                        }
                    ]
                },
            },
        )
    )


@pytest.mark.parametrize("member", ["line_chart", "table"])
@pytest.mark.parametrize(
    ("bound", "periods"),
    [
        pytest.param({}, ["2019", "2020", "2021", "2022", "2023"], id="continued"),
        pytest.param({"measurementMethod": "Estimated"}, ["2022", "2023"], id="estimates"),
    ],
)
def test_a_continued_series_exports_each_row_the_card_shows(
    member, bound, periods, frame_store, revision_store, release_registry
):
    """A series continued by its source's estimates exports its estimated rows, whether the
    card draws them with the series or alone. Issue: the Zambia chart that drew only
    2025-2026."""
    _put_continued_frame(frame_store)
    created = write_component(
        intent={
            "type": member,
            "frame_ref": _frame_ref_body(_FRAME_REF),
            "select": {"recipient": "country/ETH", **bound},
        },
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    rows = _csv_rows(_export(created, frame_store=frame_store, revision_store=revision_store).data)

    assert [row["time_period"] for row in rows] == periods
    assert [(row["series"], row["time_period"]) for row in rows] == _shown(
        created.resolved_payload["component"]
    )


def test_a_table_row_of_a_continued_series_names_its_own_method(
    frame_store, revision_store, release_registry
):
    """A table whose series differ on method labels each continuing row by its own method."""
    _put_continued_frame(frame_store)
    created = write_component(
        intent={
            "type": "table",
            "frame_ref": _frame_ref_body(_FRAME_REF),
            "select": {"stat_var": _STAT_VAR},
        },
        frame_store=frame_store,
        revision_store=revision_store,
        release_registry=release_registry,
    )

    table_rows = created.resolved_payload["component"]["rows"]
    assert [
        (row["dimensions"]["recipient"], row["period"], row["dimensions"]["measurementMethod"])
        for row in table_rows
    ] == [
        *(("Ethiopia", str(year), "Not specified") for year in range(2019, 2022)),
        *(("Ethiopia", str(year), "Estimated") for year in range(2022, 2024)),
        *(("Ghana", str(year), "Reported") for year in range(2019, 2024)),
    ]
